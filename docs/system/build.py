#!/usr/bin/env python3
"""레지스트리와 목표 OpenAPI에서 팀용 xlsx·마크다운과 실패 목록을 만든다."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SYSTEM = ROOT / "docs/system"
OUT = SYSTEM / "out"
API_OUT = ROOT / "docs/api/out"
METHODS = {"get", "post", "put", "patch", "delete", "head", "options"}
TEXT_FINGERPRINT_SUFFIXES = {".csv", ".json", ".md"}


def fingerprint_bytes(path: Path) -> bytes:
    """Hash text sources independently of the checkout's newline convention."""
    content = path.read_bytes()
    if path.suffix.lower() in TEXT_FINGERPRINT_SUFFIXES:
        return content.replace(b"\r\n", b"\n").replace(b"\r", b"\n")
    return content


def load(path: Path) -> dict:
    with path.open(encoding="utf-8") as stream:
        return yaml.safe_load(stream)


def cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return ", ".join(str(x) for x in value)
    return str(value)


def md(value) -> str:
    return cell(value).replace("|", "\\|").replace("\n", "<br>")


def table(headers: list[str], rows) -> str:
    return "| " + " | ".join(headers) + " |\n| " + " | ".join("---" for _ in headers) + " |\n" + "".join(
        "| " + " | ".join(md(x) for x in row) + " |\n" for row in rows)


def slug(value: str) -> str:
    # 원본의 잘못된 ID도 안전한 파일명으로 보존하되 정상 ID와 충돌하지 않는다.
    return "".join(c if c.isalnum() or c in "-_" else f"~{ord(c):02x}" for c in value)


def source_label(source: dict | None) -> str:
    if not source:
        return ""
    return f"{source['path']}:{source['row']} ({source['rowUnit']})"


def source_link(source: dict, depth: int = 4) -> str:
    return f"[{source_label(source)}]({'../' * depth}{source['path']})"


def operations(document: dict):
    for path, item in document.get("paths", {}).items():
        for method, op in item.items():
            if method in METHODS:
                yield path, method, op


def definition_file(definition: dict) -> str:
    return Path(definition["source"]["path"]).name


def detail_authority(specs: dict, links: dict, api: dict, requirements: dict, finding) -> dict:
    """상세 시트의 현행/이력 권한과 의미 이행을 사람이 편집한 links와 대조한다.

    목차 밖 ID, 이력 시트의 정의, 번호가 재사용된 AI ID가 번호 일치만으로 활성 목록·요구·판정에 섞이지 않게 한다.
    """
    detail = links.get("detailAuthority", {})
    historical_files = {s["file"]: s for s in detail.get("historicalSheets", [])}
    historical_only = set(detail.get("historicalOnlySpecIds", []))
    corrections = {c["rawId"]: c for c in detail.get("rawIdCorrections", [])}
    active = {rid for rid in specs if rid not in historical_only and rid not in corrections}
    ops = {op["operationId"]: (path, method, op) for path, method, op in operations(api)}

    def invalid(target, message):
        finding("DETAIL_AUTHORITY_INVALID", target, message, "error")

    current_files = {definition_file(d) for rec in specs.values() for d in rec["definitions"]}
    for sheet in historical_files.values():
        if not sheet.get("sameIdMeaning"):
            # 같은 번호가 다른 의미로 재사용될 수 있는 이력 시트는 모든 정의에 의미 이행 항목이 있어야 한다.
            migrated = {(m["legacySpecId"], m["legacySource"]["file"], m["legacySource"]["row"]) for m in detail.get("legacyAiMeaning", [])}
            for rid, rec in specs.items():
                for d in rec["definitions"]:
                    if definition_file(d) == sheet["file"] and (rid, sheet["file"], d["source"]["row"]) not in migrated:
                        invalid(rid, f"이력 시트 {sheet['file']} {d['source']['row']}행의 정의에 의미 이행 항목(legacyAiMeaning)이 없다. 번호 일치로 현재 ID에 연결하지 않는다.")
        if sheet["file"] not in current_files or sheet["supersededBy"] not in current_files:
            invalid(sheet["file"], "이력 시트 또는 대체 시트가 현재 고정본 상세에 없다.")
        if sheet.get("sourceCapture"):
            # 원본에서 삭제된 이력 탭은 선언한 이전 고정본에서만 읽어야 한다. 다른 수집본의 같은 이름 파일로 바뀌지 않게 한다.
            paths = {Path(d["source"]["path"]).parent.as_posix() for rec in specs.values() for d in rec["definitions"]
                     if definition_file(d) == sheet["file"]}
            if paths != {sheet["sourceCapture"]}:
                invalid(sheet["file"], f"이력 시트의 출처 고정본이 선언({sheet['sourceCapture']})과 다르다: {sorted(paths)}")
        for rid in sheet["specIds"]:
            files = [definition_file(d) for d in specs[rid]["definitions"]] if rid in specs else []
            if sheet["file"] not in files:
                invalid(rid, f"이력 시트 {sheet['file']}에 이 ID의 상세가 없다.")
            if rid in active and sheet["supersededBy"] not in files:
                invalid(rid, f"활성 ID에 대체 시트 {sheet['supersededBy']}의 현행 상세가 없다.")
    for rid, rec in specs.items():
        if rec["status"] == "목차 없음" and rid not in historical_only and rid not in corrections:
            invalid(rid, "목차 밖 상세 ID가 이력/정정 목록에 없다. 활성 목록으로 자동 승격하지 않는다.")
    for rid in sorted(historical_only):
        rec = specs.get(rid)
        if not rec or rec["status"] != "목차 없음":
            invalid(rid, "이력 전용 ID가 없거나 현재 목차에 존재한다.")
        elif any(definition_file(d) not in historical_files for d in rec["definitions"]):
            invalid(rid, "이력 전용 ID에 현행 시트의 상세가 있다.")
    for rid, c in corrections.items():
        rec, target = specs.get(rid), specs.get(c["effectiveSpecId"])
        if not rec or not target or c["effectiveSpecId"] not in active:
            invalid(rid, "원문 ID 정정 연결의 대상이 없거나 활성 ID가 아니다.")
            continue
        if not any(definition_file(d) == c["file"] and d["source"]["row"] == c["idRow"]
                   and (d.get("heading") or {}).get("row") == c["headingRow"]
                   and (d.get("heading") or {}).get("text", "").startswith(c["effectiveSpecId"]) for d in rec["definitions"]):
            invalid(rid, "원문 ID 행/제목 행/제목의 연결 ID가 현재 고정본과 다르다.")
    resolved = detail.get("resolvedRawIdCorrections", [])
    for c in resolved:
        target = specs.get(c["effectiveSpecId"])
        if c["rawId"] in specs or not target or c["effectiveSpecId"] not in active:
            invalid(c["rawId"], "해소로 기록한 원문 오류 ID가 현재 고정본에 남아 있거나 대상 ID가 활성 ID가 아니다.")
        elif not any(definition_file(d) == c["file"] and d["source"]["row"] == c["idRow"]
                     and (d.get("heading") or {}).get("row") == c["headingRow"] for d in target["definitions"]):
            invalid(c["rawId"], "해소로 기록한 기능 ID 행·제목 행이 현재 고정본의 대상 ID 정의와 다르다.")
    for item in detail.get("legacyAiMeaning", []):
        legacy, source = item["legacySpecId"], item["legacySource"]
        if legacy not in specs or not any(definition_file(d) == source["file"] and d["source"]["row"] == source["row"]
                                          for d in specs[legacy]["definitions"]) or source["file"] not in historical_files:
            invalid(legacy, "옛 AI 정의의 출처 시트/행이 현재 고정본 이력 정의와 다르다.")
        for target in item["currentSpecIds"] + item.get("implementationJudgmentTo", []):
            if target not in active:
                invalid(legacy, f"의미 이행 대상 {target}가 활성 ID가 아니다.")
        for qid in item.get("requirementIds", []):
            if qid not in requirements:
                invalid(legacy, f"요구사항 {qid}가 없다.")
    for binding in links.get("requirementBindings", []):
        rid = binding["specId"]
        if rid not in active or binding["rawRequirementIds"] != specs[rid]["requirementIds"]:
            invalid(rid, "요구 의미 연결의 명세가 활성 ID가 아니거나 원문 요구 번호가 목차와 다르다.")
        for qid in binding["effectiveRequirementIds"]:
            if qid not in requirements:
                invalid(rid, f"의미 기준 요구사항 {qid}가 없다.")
    for reserved in links.get("reservedRequirements", []):
        if reserved["requirementId"] not in requirements or reserved["rawIndexSpecId"] not in specs \
                or any(q not in requirements for q in reserved.get("relatedRequirementIds", [])):
            invalid(reserved["requirementId"], "유보 요구 또는 관련 요구/목차 ID가 없다.")
        for oid in reserved["operationIds"]:
            if oid not in ops or reserved["requirementId"] not in ops[oid][2]["x-zoop"].get("requirementIds", []):
                invalid(reserved["requirementId"], f"{oid}가 목표 계약의 x-zoop.requirementIds에 이 요구를 연결하지 않는다.")
    scope = links.get("authority", {}).get("implementationScope", {})
    for exception in scope.get("exceptions", []):
        for rid in exception.get("specIds", []):
            if rid not in active:
                invalid(exception["id"], f"범위 예외 명세 {rid}가 활성 ID가 아니다.")
        for qid in exception.get("requirementIds", []):
            if qid not in requirements:
                invalid(exception["id"], f"범위 예외 요구사항 {qid}가 없다.")
        for oid in exception.get("retainedOperationIds", []):
            status = ops[oid][2]["x-zoop"].get("featureStatus", {}) if oid in ops else {}
            if oid not in ops or (exception.get("specIds") and status.get("exceptionRef") != exception["id"]):
                invalid(exception["id"], f"유보 작업 {oid}가 목표 계약에 없거나 featureStatus.exceptionRef가 다르다.")
    by_id = {e["id"]: e for e in scope.get("exceptions", [])}
    reserved_ops = {oid: r for r in links.get("reservedRequirements", []) for oid in r["operationIds"]}
    # 화면 없는 작업은 사용자 기능 제외로 선언된 예외의 유보 작업이고 featureStatus가 그 예외를 가리킬 때만 허용한다.
    screenless = {oid for e in by_id.values() if e.get("userFacing") is False and e.get("specIds")
                  for oid in e.get("retainedOperationIds", [])
                  if oid in ops and ops[oid][2]["x-zoop"].get("featureStatus", {}).get("exceptionRef") == e["id"]
                  and str(ops[oid][2]["x-zoop"]["featureStatus"].get("state", "")).startswith("excluded-")}
    excluded_ops = {oid for e in by_id.values() if e.get("userFacing") is False and e.get("specIds") for oid in e.get("retainedOperationIds", [])}
    for oid, (_, _, op) in ops.items():
        x = op["x-zoop"]
        status = x.get("featureStatus")
        if status:
            exception = by_id.get(status.get("exceptionRef"))
            if not exception or oid not in exception.get("retainedOperationIds", []) or not exception.get("specIds"):
                invalid(oid, "featureStatus가 선언된 범위 예외의 유보 작업이 아니다. 임의의 상태 문자열로 제외·유보를 선언할 수 없다.")
        if oid in excluded_ops and x.get("screenIds"):
            finding("EXCLUDED_FEATURE_HAS_SCREEN", oid, "사용자 기능 제외로 선언된 유보 작업은 활성 화면 연결(screenIds)을 가질 수 없다. 과거 연결은 featureStatus.formerScreenIds에만 둔다.", "error")
        if not x.get("fsIds"):
            # 공통 품질(NFS)은 기능 명세를 대신하지 못한다. FS가 없으면 명시된 유보 요구 연결이 있어야 한다.
            reserved = reserved_ops.get(oid)
            if not reserved or reserved["requirementId"] not in x.get("requirementIds", []) or reserved["requirementId"] not in requirements:
                finding("API_WITHOUT_FUNCTION", oid, "기능 명세(FS)가 없는 작업은 명시된 유보 요구 연결(reservedRequirements 작업 + 유효한 x-zoop.requirementIds)이 있어야 한다. NFS만으로는 허용되지 않는다.", "error")
    chart = links.get("authority", {}).get("chartResolution", {})
    series = api["components"]["schemas"]["PriceSeries"]["properties"]
    if chart != api.get("x-zoop-approvals", {}).get("chartResolution") \
            or set(chart.values()) != set(series["interval"].get("enum", [])) or set(chart) != set(series["period"]["enum"]):
        invalid("chartResolution", "고정 봉 해상도가 links.authority, OpenAPI 승인 기록, PriceSeries 열거값과 일치하지 않는다.")
    return {"historicalFiles": historical_files, "historicalOnly": historical_only, "corrections": corrections,
            "resolvedCorrections": resolved, "active": active, "legacyAi": detail.get("legacyAiMeaning", []),
            "ops": ops, "screenlessAllowed": screenless}


def effective_mapping(rid: str, rec: dict, definition: dict | None, ctx: dict, req_bindings: dict) -> tuple[list[str], list[str]]:
    """상세 정의 하나가 가리키는 현행 명세 ID들과 의미 기준 요구 번호를 구한다.

    이력 정의는 현재 ID의 요구 연결을 상속하지 않고 legacySource(파일·행)로 선언된 의미 이행 항목만 따른다.
    한 옛 정의가 여러 현행 ID(예: 옛 FS046의 서버 공통 의무)로 이행될 수 있다.
    """
    historical_files = ctx["historicalFiles"]
    if definition is not None and definition_file(definition) in historical_files:
        for item in ctx["legacyAi"]:
            source = item["legacySource"]
            if item["legacySpecId"] == rid and source["file"] == definition_file(definition) and source["row"] == definition["source"]["row"]:
                return list(item["currentSpecIds"]), list(item["requirementIds"])
        return [rid], list(rec["requirementIds"])
    if rid in ctx["corrections"]:
        return [ctx["corrections"][rid]["effectiveSpecId"]], list(rec["requirementIds"])
    binding = req_bindings.get(rid)
    return [rid], list(binding["effectiveRequirementIds"] if binding else rec["requirementIds"])


def spec_mapping(rid: str, rec: dict, ctx: dict, req_bindings: dict) -> tuple[list[str], list[str]]:
    """명세 ID 단위의 현행 대응. 이력 전용 ID는 선언된 모든 이력 정의의 이행 대상을 합친다."""
    if rid in ctx["historicalOnly"]:
        ids, reqs = [], []
        for definition in rec["definitions"]:
            targets, requirements = effective_mapping(rid, rec, definition, ctx, req_bindings)
            ids += targets
            reqs += requirements
        return list(dict.fromkeys(ids)), list(dict.fromkeys(reqs))
    return effective_mapping(rid, rec, None, ctx, req_bindings)


def deref(document: dict, value: dict) -> dict:
    target = value.get("$ref", "")
    if not target.startswith("#/"):
        return value
    current = document
    for part in target[2:].split("/"):
        current = current[part.replace("~1", "/").replace("~0", "~")]
    return current


def schema_label(value: dict) -> str:
    if "$ref" in value:
        return value["$ref"].rsplit("/", 1)[-1]
    union = value.get("oneOf", value.get("anyOf", []))
    if union:
        return " / ".join(schema_label(v) for v in union)
    kind = value.get("type", "object")
    if kind == "array":
        label = f"{schema_label(value.get('items', {}))}[]"
        return label + (" (uniqueItems)" if value.get("uniqueItems") else "")
    text = cell(kind)
    if "enum" in value:
        text += " {" + cell(value["enum"]) + "}"
    if "format" in value:
        text += f" ({value['format']})"
    return text


def schema_md(value: dict) -> str:
    if "$ref" in value:
        name = value["$ref"].rsplit("/", 1)[-1]
        return f"[{name}](#schema-{slug(name).lower()})"
    union = value.get("oneOf", value.get("anyOf", value.get("allOf", [])))
    if union:
        return " / ".join(schema_md(v) for v in union)
    if value.get("type") == "array":
        label = schema_md(value.get("items", {})) + "[]"
        return label + (" (uniqueItems)" if value.get("uniqueItems") else "")
    return schema_label(value)


def inline_schema_rows(schema: dict, prefix: str = "") -> list[list[str]]:
    result = []
    for key, value in schema.get("properties", {}).items():
        path = f"{prefix}.{key}" if prefix else key
        result.append([path, "필수" if key in schema.get("required", []) else "선택", schema_md(value), value.get("description", "")])
        result.extend(inline_schema_rows(value, path))
        if "items" in value:
            result.extend(inline_schema_rows(value["items"], path + "[]"))
        for union in ("oneOf", "anyOf", "allOf"):
            for branch in value.get(union, []):
                result.extend(inline_schema_rows(branch, path))
    return result


def schema_fields(document: dict, schema: dict, prefix: str = "", seen: frozenset = frozenset()) -> dict:
    """참조·배열·union을 따라 필드별 형식/필수 여부를 비교한다. 순환 참조는 한 번만 펼친다."""
    result = {}
    if "$ref" in schema:
        target = schema["$ref"]
        if target in seen:
            return {prefix or "$": {"recursive": target.rsplit("/", 1)[-1]}}
        return schema_fields(document, deref(document, schema), prefix, seen | {target})
    signature = {k: schema[k] for k in ("type", "enum", "const", "format", "minimum", "maximum", "minLength", "maxLength", "minItems", "maxItems", "uniqueItems") if k in schema}
    if signature:
        result[prefix or "$"] = signature
    required = set(schema.get("required", []))
    for name, child in schema.get("properties", {}).items():
        key = f"{prefix}.{name}" if prefix else name
        result[key + ":required"] = name in required
        result.update(schema_fields(document, child, key, seen))
    if "items" in schema:
        result.update(schema_fields(document, schema["items"], prefix + "[]", seen))
    for union in ("oneOf", "anyOf", "allOf"):
        for i, branch in enumerate(schema.get(union, [])):
            result.update(schema_fields(document, branch, f"{prefix}<{union}:{i}>", seen))
    return result


def operation_fields(document: dict, path: str, method: str, op: dict) -> dict:
    fields = {}
    params = document["paths"][path].get("parameters", []) + op.get("parameters", [])
    for parameter in params:
        p = deref(document, parameter)
        key = f"{p['in']}.{p['name']}"
        fields[key + ":required"] = p.get("required", False)
        fields.update(schema_fields(document, p.get("schema", {}), key))
    body = deref(document, op.get("requestBody", {}))
    if body:
        fields["requestBody:required"] = body.get("required", False)
    for content_type, value in body.get("content", {}).items():
        fields.update(schema_fields(document, value.get("schema", {}), "request:" + content_type))
    for status, response in op.get("responses", {}).items():
        response = deref(document, response)
        fields["response:" + status] = True
        for content_type, value in response.get("content", {}).items():
            fields.update(schema_fields(document, value.get("schema", {}), f"response:{status}:{content_type}"))
        for header, value in response.get("headers", {}).items():
            fields.update(schema_fields(document, deref(document, value).get("schema", {}), f"response:{status}:header:{header}"))
    fields["security"] = op.get("security", document.get("security", []))
    return fields


def workbook(path: Path, sheets: dict[str, tuple[list[str], list[list]]]) -> dict:
    book = Workbook()
    book.remove(book.active)
    for name, (headers, rows) in sheets.items():
        ws = book.create_sheet(name)
        ws.append(headers)
        for row in rows:
            # 원본 셀을 수식으로 실행하지 않는다(구글 시트 가져오기에서도 문자열이다).
            ws.append([cell(v) for v in row])
            for entry in ws[ws.max_row]:
                if entry.data_type == "f":
                    entry.data_type = "s"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for entry in ws[1]:
            entry.font = Font(bold=True, color="FFFFFF")
            entry.fill = PatternFill("solid", fgColor="17463E")
        for row in ws.iter_rows(min_row=2):
            for entry in row:
                entry.alignment = Alignment(vertical="top", wrap_text=True)
        for index, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(index)].width = 48 if any(k in header for k in ("사유", "내용", "경고", "차이", "출처", "필요", "요청", "응답")) else 26
    book.save(path)
    # 저장 후 독립 reader로 전체 행/셀을 되읽어 배포본과 생성 입력을 대조한다.
    readback = load_workbook(path, read_only=True, data_only=False)
    try:
        if readback.sheetnames != list(sheets):
            raise ValueError("xlsx 되읽기: 시트 목록/순서 불일치")
        counts = {}
        for name, (headers, rows) in sheets.items():
            expected = [headers] + rows
            actual = list(readback[name].iter_rows(values_only=True))
            if len(actual) != len(expected):
                raise ValueError(f"xlsx 되읽기: {name} 행 수 불일치")
            for n, (saved, original) in enumerate(zip(actual, expected), 1):
                if tuple(cell(v) for v in saved) != tuple(cell(v) for v in original):
                    raise ValueError(f"xlsx 되읽기: {name} {n}행 셀 불일치")
            counts[name] = len(rows)
        return {"status": "passed", "sheetDataRows": counts,
                "cells": sum((len(rows) + 1) * len(headers) for headers, rows in sheets.values())}
    finally:
        readback.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="생성 후 실패 목록을 출력한다. 목록이 하나라도 있으면 종료 코드 1")
    parser.add_argument("--bff-url", help="현재 BFF의 /openapi.json URL. 지정 시 해당 문서만 GET하고 out에 관측본을 저장한다")
    args = parser.parse_args()
    catalog = load(SYSTEM / "catalog.yaml")
    links = load(SYSTEM / "links.yaml")
    api = load(ROOT / links["apiContract"])
    OUT.mkdir(parents=True, exist_ok=True)
    API_OUT.mkdir(parents=True, exist_ok=True)
    for pattern in ("screen-*.md", "feature-*.md", "api-differences/*.json"):
        # 이 생성기가 관리하는 산출물만 제거하여 사라진 ID/작업의 문서가 남지 않게 한다.
        for stale in OUT.glob(pattern):
            stale.unlink()
    findings: list[dict] = []

    def finding(code, target, message, severity="warning", source=""):
        findings.append({"code": code, "target": target, "severity": severity, "message": message, "source": source})

    catalog_warnings = catalog["warnings"]
    for path, digest in catalog["fingerprints"].items():
        actual = ROOT / path
        if not actual.exists() or hashlib.sha256(fingerprint_bytes(actual)).hexdigest() != digest:
            finding("SOURCE_CHANGED", path, "가져오기 이후 원본이 바뀌었다. 동일 명령으로 다시 가져와야 한다.", "error")
    specs = {r["id"]: r for r in catalog["specs"]}
    screens = {r["id"]: r for r in catalog["screens"]}
    tasks = {r["id"]: r for r in catalog["wbs"]}
    requirements = defaultdict(list)
    for rec in catalog["requirements"]:
        requirements[rec["id"]].append(rec)
    implementations = {r["specId"]: r for r in catalog["implementation"]}
    decisions = links.get("featureDecisions", {})
    authority = links.get("authority", {})
    scope = authority.get("implementationScope", {})
    raw_proposal_ids = {rid for rid, rec in specs.items() if rec.get("scopeClass") == "proposal"}
    included_proposals = scope.get("includeRawProposalIds", [])
    exceptions = scope.get("exceptions", [])
    exception_specs = {rid: e for e in exceptions for rid in e.get("specIds", [])}
    authority_source = ROOT / authority.get("source", "")
    authority_valid = (
        bool(authority.get("decisionId")) and bool(authority.get("authority"))
        and authority_source.is_file()
        and scope.get("appliesTo") == "all-specs-and-wbs-except-recorded-exceptions"
        and scope.get("status") == "팀합의요구/설계대상"
        and scope.get("implementationComplete") is False
        and scope.get("paidExecutionApproved") is False
        and len(included_proposals) == len(set(included_proposals))
        and set(included_proposals) == raw_proposal_ids
        and len({e["id"] for e in exceptions}) == len(exceptions)
    )
    if not authority_valid:
        finding("INVALID_EFFECTIVE_AUTHORITY", "authority", "유효 사용자 결정의 출처·전 범위(기록된 예외 제외)·원문 제안 ID·완료/실행 경계를 확인해야 한다.", "error")
    api_authority = api.get("x-zoop-approvals", {}).get("implementationScope", {})
    if api_authority != {"authorityRef": "docs/system/links.yaml#authority",
                         "decisionId": authority.get("decisionId"), "source": authority.get("source")}:
        finding("API_AUTHORITY_MISMATCH", "OpenAPI", "목표 API의 유효 기준 참조가 links.yaml.authority와 다르다.", "error")
    effective_scope = scope["status"] if authority_valid else "유효 기준 오류"
    detail_ctx = detail_authority(specs, links, api, requirements, finding)
    active_ids, historical_only = detail_ctx["active"], detail_ctx["historicalOnly"]
    raw_corrections, historical_files = detail_ctx["corrections"], detail_ctx["historicalFiles"]
    legacy_ai = detail_ctx["legacyAi"]
    req_bindings = {b["specId"]: b for b in links.get("requirementBindings", [])}
    reserved_requirements = links.get("reservedRequirements", [])

    def historical_definitions(rec):
        return [d for d in rec["definitions"] if definition_file(d) in historical_files]

    def current_definitions(rec):
        return [d for d in rec["definitions"] if definition_file(d) not in historical_files]

    def scope_state(rid):
        if rid in historical_only:
            return "historical-only"
        if rid in raw_corrections:
            return "raw-malformed-id"
        return "excluded-retained" if rid in exception_specs else "active"

    def spec_scope(rid):
        state = scope_state(rid)
        if state == "historical-only":
            return "이력 상세(현행 목차 밖) · 현재 의미는 의미 이행표 참조"
        if state == "raw-malformed-id":
            return f"원문 ID 오류 · {raw_corrections[rid]['effectiveSpecId']}에 유효 연결"
        return exception_specs[rid]["effectiveStatus"] if state == "excluded-retained" else effective_scope

    legacy_by_id = {m["legacySpecId"]: m for m in legacy_ai}
    judgments_to = defaultdict(list)
    for m in legacy_ai:
        if m["legacySpecId"] in implementations:
            for target in m.get("implementationJudgmentTo", []):
                judgments_to[target].append({**implementations[m["legacySpecId"]], "legacySpecId": m["legacySpecId"]})

    def implementation_records(rid, historical_definition=False):
        """번호가 재사용된 AI ID의 옛 판정을 같은 번호의 현재 정의에 붙이지 않는다. 의미 이행표가 정한 곳에만 적용한다."""
        if historical_definition or rid not in active_ids:
            return [implementations[rid]] if rid in implementations else []
        if rid in judgments_to:
            return judgments_to[rid]
        return [] if rid in legacy_by_id else ([implementations[rid]] if rid in implementations else [])

    def implementation_text(records, key, rid):
        values = [(r.get("legacySpecId", rid), r.get(key)) for r in records]
        if not values:
            return ""
        if len(values) == 1 and values[0][0] == rid:
            return cell(values[0][1]) if key != "source" else source_label(values[0][1])
        return " / ".join(f"옛 {who}: {source_label(v) if key == 'source' else v}" for who, v in values)

    for w in catalog_warnings:
        record = specs.get(w["id"])
        if w["code"] == "MULTIPLE_SPEC_DEFINITIONS" and record and historical_definitions(record) and len(current_definitions(record)) == 1:
            continue  # 같은 ID의 옛 시트 정의는 아래 HISTORICAL_DETAIL_SUPERSEDED로 구분해 남긴다.
        finding(w["code"], w["id"], w["message"], source=source_label(w.get("source")))
    for rid in specs:
        record = specs[rid]
        if rid in active_ids and historical_definitions(record) and len(current_definitions(record)) == 1:
            sheet = historical_files[definition_file(historical_definitions(record)[0])]
            finding("HISTORICAL_DETAIL_SUPERSEDED", rid, f"옛 시트 {sheet['file']}의 정의는 이력으로 보존한다. 현행 상세는 {sheet['supersededBy']}이며 같은 번호라도 의미가 다를 수 있어 번호 일치로 합치지 않는다.", source=source_label(historical_definitions(record)[0]["source"]))
    for rid in sorted(historical_only):
        current = sorted({t for m in legacy_ai if m["legacySpecId"] == rid for t in m["currentSpecIds"]})
        finding("HISTORICAL_ONLY_ID", rid, f"현재 목차 밖의 이력 ID다. 활성 목록으로 승격하지 않고 현재 의미는 {', '.join(current) or '미연결'}로 이행한다.", source=source_label(specs[rid]["source"]))
    for rid, correction_link in raw_corrections.items():
        finding("RAW_ID_CORRECTION", rid, f"원문 기능 ID 셀({correction_link['file']} {correction_link['idRow']}행)은 정정되지 않았다. 제목 행({correction_link['headingRow']}행)과 목차 통합 선언에 근거해 {correction_link['effectiveSpecId']}에 유효 연결만 기록한다.", source=source_label(specs[rid]["source"]))
    for exception in exceptions:
        finding("SCOPE_EXCEPTION", exception["id"], f"{exception['effectiveStatus']}: {exception['reason']}", source=exception.get("basis", ""))
    for reserved in reserved_requirements:
        finding("RESERVED_REQUIREMENT_UNMAPPED", f"요구{reserved['requirementId']}", f"{reserved['effectiveStatus']}: {reserved['reason']}", source=source_label(requirements[reserved["requirementId"]][0]["source"]) if reserved["requirementId"] in requirements else "")
    for rid in (r for r in specs if r in active_ids):
        record = specs[rid]
        match = re.fullmatch(r"A(\d+):B\d+", record["detailRange"].strip())
        rows = [(d.get("heading") or {}).get("row") for d in current_definitions(record)]
        if match and rows and int(match[1]) not in rows:
            finding("INDEX_DETAIL_RANGE_STALE", rid, f"목차 상세 위치 {record['detailRange']}와 현행 상세 구획 제목 행 {rows}이 다르다. 원문 값을 보존하며 자동 정정하지 않는다.", source=source_label(record["source"]))
        for definition in current_definitions(record):
            meta = re.search(r"요구사항:\s*([\d.,\s]+)", definition.get("meta", ""))
            numbers = sorted({x.strip().removesuffix(".0") for x in meta[1].split(",") if x.strip()}) if meta else []
            if numbers and numbers != sorted(record["requirementIds"]):
                finding("INDEX_REQUIREMENT_META_MISMATCH", rid, f"목차 요구 번호 {record['requirementIds']}와 현행 상세 메타 요구 번호 {numbers}가 다르다. 원문을 보존한다.", source=source_label(definition["source"]))
    effective_specs = {
        rid: {"raw": {"status": rec["status"], "priority": rec["priority"],
                      "scopeClass": rec.get("scopeClass", "existing"), "source": rec["source"]},
              "implementationScope": spec_scope(rid), "scopeState": scope_state(rid),
              "effectiveSpecIds": spec_mapping(rid, rec, detail_ctx, req_bindings)[0],
              "effectiveRequirementIds": spec_mapping(rid, rec, detail_ctx, req_bindings)[1],
              "authorityDecisionId": authority.get("decisionId"),
              "implementationEvidence": "; ".join(r.get("evidenceKind", "") for r in implementation_records(rid)) or "문서 판정 없음"}
        for rid, rec in specs.items()
    }
    wbs_bindings = {r["wbsId"]: r for r in links.get("wbsConnections", [])}
    common_wbs = links.get("commonWbs", {})
    for wid, title in common_wbs.items():
        if wid not in tasks or tasks[wid]["title"] != title:
            finding("WBS_MEANING_CHANGED", wid, f"공통 계약 작업 '{title}'과 현재 WBS 의미가 다르다.", "error")
    by_spec, by_screen, by_wbs = defaultdict(list), defaultdict(list), defaultdict(list)
    all_ops = list(operations(api))
    seen_ops = set()
    for path, method, op in all_ops:
        oid = op["operationId"]
        if oid in seen_ops:
            finding("DUPLICATE_OPERATION_ID", oid, "목표 API operationId 중복", "error")
        seen_ops.add(oid)
        x = op.get("x-zoop", {})
        for key, valid, index in (("fsIds", specs, by_spec), ("nfsIds", specs, by_spec), ("screenIds", screens, by_screen), ("wbsIds", tasks, by_wbs)):
            for rid in x.get(key, []):
                if rid not in valid:
                    finding("BROKEN_ID", oid, f"{key}: {rid} 참조 대상이 없다.", "error")
                index[rid].append((path, method, op))
        if not x.get("screenIds") and oid not in detail_ctx["screenlessAllowed"]:
            finding("API_WITHOUT_SCREEN", oid, "목표 API에 연결 화면이 없다.", "error")
        if x.get("priority") not in ("P1", "P2"):
            finding("INVALID_PRIORITY", oid, "API 우선순위는 P1 또는 P2여야 한다.", "error")
        if x.get("teamQuestions"):
            finding("TEAM_DECISION_REQUIRED", oid, " / ".join(x["teamQuestions"]))
        for code in x.get("errorCodes", []):
            if code not in {e["code"] for e in api["x-zoop-error-codes"]}:
                finding("BROKEN_ERROR_CODE", oid, f"정의되지 않은 오류 코드 {code}", "error")
        for wid in x.get("wbsIds", []):
            if wid not in tasks or wid in common_wbs:
                continue
            binding = wbs_bindings.get(wid)
            if not binding or not set(binding["specIds"]).intersection(x.get("fsIds", []) + x.get("nfsIds", [])):
                finding("API_WBS_MEANING_MISMATCH", oid, f"{wid}: 현재 제목/역할 기반 편집 연결과 API 기능 의미가 대응하지 않는다.", "error")
        for rid in x.get("fsIds", []) + x.get("nfsIds", []):
            if rid in specs and rid not in active_ids:
                finding("API_LINKS_NON_ACTIVE_ID", oid, f"{rid}: 이력 전용/원문 오류 ID는 활성 목표 작업에 연결하지 않는다. 의미가 이행된 현재 ID로 연결해야 한다.", "error")
            if not authority_valid:
                finding("API_SCOPE_WITHOUT_AUTHORITY", oid, f"{rid}: 유효 기준이 검증되지 않았다.", "error")
            if decisions.get(rid, {}).get("apiDisposition") == "API 불필요":
                finding("OFFLINE_FEATURE_HAS_API", oid, f"{rid}: API 불필요 판정과 HTTP 작업 연결이 충돌한다.", "error")
        for qid in x.get("requirementIds", []):
            if qid not in requirements:
                finding("BROKEN_ID", oid, f"requirementIds: {qid} 요구사항 원본 행이 없다.", "error")
        if "FS042" in x.get("fsIds", []) and not x.get("meaningQualifications", {}).get("FS042"):
            finding("SPEC_MEANING_UNQUALIFIED", oid, "FS042의 시점 그래프/기존 재무 중 연결 의미를 명시해야 한다.", "error")
    for rid, rec in specs.items():
        for qid in rec["requirementIds"]:
            if qid not in requirements:
                finding("BROKEN_REQUIREMENT_ID", rid, f"요구사항 {qid} 원본 행이 없다.", "error", source_label(rec["source"]))
            elif len(requirements[qid]) == 1 and requirements[qid][0]["title"] != rec["title"]:
                binding = req_bindings.get(rid)
                if binding:
                    finding("REQUIREMENT_BINDING_BY_MEANING", rid, f"명세 '{rec['title']}' ↔ 원문 요구 번호 {qid} '{requirements[qid][0]['title']}'. 의미 기준 연결은 요구 {binding['effectiveRequirementIds'] or '없음'}이다({binding['relation']}): {binding['note']}", source=source_label(requirements[qid][0]["source"]))
                else:
                    finding("REQUIREMENT_TITLE_REVIEW", rid, f"명세 '{rec['title']}' ↔ 요구사항 {qid} '{requirements[qid][0]['title']}'. 분리·통합/번호 이동인지 확인 필요(자동 재연결하지 않음).", source=source_label(requirements[qid][0]["source"]))
        disposition = decisions.get(rid, {})
        if rec.get("scopeClass") == "proposal":
            finding("RAW_SCOPE_STATUS_CORRECTION", rid, f"원문 '{rec['status']}'/우선순위 '{rec['priority']}'를 보존한다. 현재 {authority.get('decisionId')}의 유효 범위는 '{effective_scope}'이며 WBS 표기 충돌은 정정 요청이다. 구현 완료·유료 실행 승인은 별개다.", source=source_label(rec["source"]))
            if disposition.get("rawScopeStatus") != rec["status"]:
                finding("RAW_SCOPE_STATUS_MISMATCH", rid, "사람 편집 rawScopeStatus가 원문 상태와 다르다.", "error")
        if rid in active_ids and ("P1" in rec["priority"] or rid in raw_proposal_ids) and not by_spec[rid] and disposition.get("apiDisposition") not in ("계약 미정", "API 불필요"):
            finding("P1_WITHOUT_API", rid, "API 대응이 필요한 P1/유효 포함 요구에 작업 연결 또는 계약 미정 사유가 없다.", "error")
        if disposition.get("apiDisposition") in ("계약 미정", "API 불필요") and not disposition.get("reason"):
            finding("API_DISPOSITION_WITHOUT_REASON", rid, "API 불필요/계약 미정에는 사유가 필요하다.", "error")
        if disposition.get("apiDisposition") == "계약 미정":
            finding("CONTRACT_UNDECIDED", rid, disposition.get("reason", "계약 미정 사유가 없다."))
    for rid, decision in decisions.items():
        if rid not in specs:
            finding("BROKEN_ID", rid, "사람 편집 featureDecisions에 존재하지 않는 명세 ID가 있다.", "error")
        for key, valid in (("screenIds", screens), ("wbsIds", tasks)):
            for target in decision.get(key, []):
                if target not in valid:
                    finding("BROKEN_ID", rid, f"{key}: {target}", "error")
        for target in decision.get("reusesSpecIds", []):
            if target not in specs:
                finding("BROKEN_ID", rid, f"재사용 근거 명세 {target}가 없다.", "error")
    human_tasks = defaultdict(list)
    for connection in links.get("wbsConnections", []):
        wid = connection["wbsId"]
        if wid not in tasks:
            finding("BROKEN_ID", wid, "사람 편집 WBS 연결의 작업 번호가 없다.", "error")
        elif connection.get("title") != tasks[wid]["title"] or connection.get("role") != tasks[wid].get("role"):
            finding("WBS_MEANING_CHANGED", wid, "번호는 존재하지만 사람이 확인한 제목/역할과 새 원문이 다르다. 내용 기반으로 다시 연결해야 한다.", "error")
        for rid in connection["specIds"]:
            if rid not in specs:
                finding("BROKEN_ID", wid, f"사람 편집 WBS 연결 명세 {rid}가 없다.", "error")
            elif rid not in active_ids:
                finding("WBS_BINDING_NON_ACTIVE_ID", wid, f"사람 편집 WBS 연결이 이력 전용/원문 오류 ID {rid}를 가리킨다. 의미가 이행된 현재 ID로 연결해야 한다.", "error")
            human_tasks[rid].append(wid)
    for wid, task in tasks.items():
        if re.fullmatch(r"\d{4}-\d\d-\d\d", task["start"]) and re.fullmatch(r"\d{4}-\d\d-\d\d", task["end"]) and task["start"] > task["end"]:
            finding("WBS_DATE_ORDER_INVALID", wid, f"원문 시작일 {task['start']}이 종료일 {task['end']}보다 늦다. 날짜를 추정·수정하지 않고 일정 정정 대상으로 보존한다.", source=source_label(task["source"]))
        for raw in task["rawSpecIds"]:
            if raw not in specs:
                finding("WBS_RAW_ID_REVIEW", wid, f"원본 WBS의 {raw}는 현재 명세 ID가 아니다. 사람이 정한 연결과 구분하여 원문을 보존한다.", source=source_label(task["source"]))
            elif raw not in wbs_bindings.get(wid, {}).get("specIds", []):
                finding("WBS_EDITORIAL_ID_DIFFERENCE", wid, f"원문 {raw}는 의존/재사용 참조 또는 업무명과 다른 산출물 ID다. 제목·역할 기반 편집 연결과 rawSpecIds를 분리하며 원문을 자동 치환하지 않는다.", source=source_label(task["source"]))
    code_by_spec = defaultdict(list)
    for row in links.get("implementationLocations", []):
        if row["specId"] not in specs:
            finding("BROKEN_ID", row["specId"], "구현 위치에 존재하지 않는 명세 ID가 있다.", "error")
        elif row["specId"] not in active_ids:
            finding("CODE_LOCATION_NON_ACTIVE_ID", row["specId"], "구현 위치가 이력 전용/원문 오류 ID에 연결됐다.", "error")
        code_by_spec[row["specId"]].extend(row["codeLocations"])
        for location in row["codeLocations"]:
            if not (ROOT / location.split(":", 1)[0]).exists():
                finding("CODE_LOCATION_MISSING", row["specId"], location, "error")

    observed_path = OUT / "bff-openapi.json"
    observation_meta = OUT / "bff-observation.json"
    bff = None
    if args.bff_url:
        try:
            with urllib.request.urlopen(args.bff_url, timeout=15) as response:
                bff = json.load(response)
            if not isinstance(bff, dict) or "paths" not in bff:
                raise ValueError("OpenAPI paths가 없다")
            observed_path.write_text(json.dumps(bff, ensure_ascii=False, indent=2) + "\n")
            observation_meta.write_text(json.dumps({"url": args.bff_url, "observedAt": datetime.now(timezone.utc).isoformat()}, ensure_ascii=False, indent=2) + "\n")
        except (OSError, ValueError, urllib.error.URLError) as error:
            finding("BFF_UNOBSERVED", args.bff_url, f"현재 BFF 관측 실패: {type(error).__name__}. 이전 관측본으로 성공 판정하지 않는다.", "error")
    elif observed_path.exists():
        bff = json.loads(observed_path.read_text())
        finding("BFF_CACHED_OBSERVATION", "BFF", "현재 서버를 조회하지 않고 out/bff-openapi.json의 마지막 관측본과 비교했다. 최종 통합은 --bff-url로 다시 실행한다.")
    else:
        finding("BFF_UNOBSERVED", "BFF", "--bff-url을 지정하지 않았고 저장된 관측본도 없다.", "error")
    differences = []
    if bff is not None:
        used = set()
        for path, method, op in all_ops:
            oid, proto = op["operationId"], op["x-zoop"]["prototype"]
            bpath = proto.get("path")
            bmethod = (proto.get("method") or method).lower()
            if not bpath:
                differences.append([oid, path, "없음", "직접 대응 BFF 작업 없음", "목표 API만 있음"])
                finding("TARGET_ONLY_API", oid, f"{method.upper()} {path}: 직접 대응 BFF 없음", "difference")
                continue
            used.add((bpath, bmethod))
            if bpath not in bff["paths"] or bmethod not in bff["paths"][bpath]:
                finding("BFF_OPERATION_MISSING", oid, f"{bmethod.upper()} {bpath}", "error")
                continue
            boperation = bff["paths"][bpath][bmethod]
            if proto.get("operationId") != boperation.get("operationId"):
                finding("BFF_OPERATION_ID_CHANGED", oid, f"기록 {proto.get('operationId')} ↔ 현재 {boperation.get('operationId')}", "difference")
            goal_fields = operation_fields(api, path, method, op)
            actual_fields = operation_fields(bff, bpath, bmethod, boperation)
            added = sorted(goal_fields.keys() - actual_fields.keys())
            removed = sorted(actual_fields.keys() - goal_fields.keys())
            changed = sorted(k for k in goal_fields.keys() & actual_fields.keys() if goal_fields[k] != actual_fields[k])
            description = f"목표만 {len(added)} · BFF만 {len(removed)} · 형식/필수/권한 변경 {len(changed)}"
            differences.append([oid, path, bpath, description, " / ".join(proto.get("differences", []))])
            if path != bpath:
                finding("BFF_PATH_DIFFERENCE", oid, f"목표 {path} ↔ BFF {bpath}", "difference")
            if added or removed or changed:
                finding("BFF_FIELD_DIFFERENCE", oid, description, "difference")
            # 개별 필드 이름과 양쪽의 실제 스키마 서명을 생략 없이 보존한다.
            (OUT / "api-differences").mkdir(exist_ok=True)
            (OUT / "api-differences" / (slug(oid) + ".json")).write_text(json.dumps({"operationId": oid, "targetPath": path, "bffPath": bpath,
                "targetOnly": {k: goal_fields[k] for k in added}, "bffOnly": {k: actual_fields[k] for k in removed},
                "changed": {k: {"target": goal_fields[k], "bff": actual_fields[k]} for k in changed}}, ensure_ascii=False, indent=2) + "\n")
        excluded = {(r["path"], r["method"].lower()): r["reason"] for r in links.get("prototypeOnly", [])}
        for path, method, op in operations(bff):
            if (path, method) not in used:
                reason = excluded.get((path, method))
                finding("PROTOTYPE_ONLY_API" if reason else "BFF_UNMAPPED_API", op["operationId"], f"{method.upper()} {path}: {reason or '대응 목표 계약이 없다. 팀 확인 필요.'}", "difference")
        for path, method in excluded:
            if path not in bff["paths"] or method not in bff["paths"][path]:
                finding("STALE_PROTOTYPE_EXCLUSION", path, "프로토타입 전용 목록에서 대응 경로가 사라졌다.")

    acceptance_rows, acceptance_records = [], []
    for rid, rec in specs.items():
        definitions = rec["definitions"]
        # 목차에만 있는 ID도 한 행으로 남겨 상세 누락을 coverage에서 숨기지 않는다.
        for index, definition in enumerate(definitions or [{"source": rec["source"], "fields": {}}], 1):
            fields = definition["fields"]
            historical = bool(definitions) and definition_file(definition) in historical_files
            effective_ids, effective_requirements = effective_mapping(rid, rec, definition if definitions else None, detail_ctx, req_bindings)
            impls = implementation_records(rid, historical)
            raw_acceptance = fields.get("수용 기준", "")
            if historical:
                detail_state = f"이력 정의(대체 {historical_files[definition_file(definition)]['supersededBy']})"
                scope_text = f"이력 정의 · 현행 범위 아님({spec_scope(rid)})" if rid in active_ids else spec_scope(rid)
            else:
                detail_state = "현행 정의" + ("(원문 ID 오류, 정정 연결)" if rid in raw_corrections else "")
                scope_text = spec_scope(rid)
            record = {
                "specId": rid, "definitionIndex": index if definitions else None,
                "indexTitle": rec["title"],
                "definitionTitle": fields.get("기능명", fields.get("요구사항명", "")),
                "definitionSource": definition["source"], "indexSource": rec["source"],
                "rawAcceptance": raw_acceptance,
                "acceptancePresence": "있음" if raw_acceptance else "없음" if definitions else "상세 정의 없음",
                "rawStatus": rec["status"], "rawPriority": rec["priority"],
                "detailAuthority": "historical" if historical else "current", "detailAuthorityLabel": detail_state,
                "effectiveSpecIds": effective_ids, "effectiveRequirementIds": effective_requirements,
                "implementationScope": scope_text, "authorityDecisionId": authority.get("decisionId"),
                "legacyImplementation": impls,
                "legacyQualification": "ID 단위 기존 문서 판정이며 이 상세 정의의 수용 기준 통과 판정이 아님",
                "codeLocations": list(dict.fromkeys([c for r in impls for c in r.get("codeLocations", [])] + (code_by_spec[rid] if not historical and rid in active_ids else []))),
                "currentRuntimeVerification": "미실행",
            }
            acceptance_records.append(record)
            acceptance_rows.append([
                rid, record["definitionIndex"], rec["title"], record["definitionTitle"],
                source_label(definition["source"]), record["acceptancePresence"],
                raw_acceptance or record["acceptancePresence"], rec["priority"], rec["status"],
                scope_text, authority.get("decisionId"), implementation_text(impls, "status", rid) or "문서 판정 없음",
                implementation_text(impls, "source", rid), record["legacyQualification"],
                implementation_text(impls, "evidenceKind", rid) or "문서 판정 없음", record["codeLocations"], "미실행",
                detail_state, effective_ids, effective_requirements,
            ])
    expected_definitions = sum(len(rec["definitions"]) for rec in specs.values())
    actual_definitions = sum(row["definitionIndex"] is not None for row in acceptance_records)
    expected_sources = Counter((rid, d["source"]["path"], d["source"]["row"])
                               for rid, rec in specs.items() for d in rec["definitions"])
    actual_sources = Counter((r["specId"], r["definitionSource"]["path"], r["definitionSource"]["row"])
                             for r in acceptance_records if r["definitionIndex"] is not None)
    acceptance_coverage = {
        "status": "passed" if set(r["specId"] for r in acceptance_records) == set(specs)
                  and expected_sources == actual_sources and actual_definitions == expected_definitions else "failed",
        "specIds": len(specs), "activeIndexSpecIds": len(active_ids), "coveredSpecIds": len({r["specId"] for r in acceptance_records}),
        "expectedDefinitions": expected_definitions, "coveredDefinitions": actual_definitions,
        "currentDefinitions": sum(r["detailAuthority"] == "current" and r["definitionIndex"] is not None for r in acceptance_records),
        "historicalDefinitions": sum(r["detailAuthority"] == "historical" for r in acceptance_records),
        "rows": len(acceptance_records),
        "missingAcceptanceDefinitions": sum(r["definitionIndex"] is not None and not r["rawAcceptance"] for r in acceptance_records),
        "missingDetailIds": [rid for rid, rec in specs.items() if not rec["definitions"]],
        "currentRuntimeVerification": "미실행",
    }
    if acceptance_coverage["status"] != "passed":
        finding("ACCEPTANCE_COVERAGE_INCOMPLETE", "수용 기준", "전체 ID/상세 정의 출처별 수용 기준 행이 빠졌다.", "error")

    corrections = []

    def correction(code, target, left, right, decision, request):
        corrections.append({"code": code, "target": target, "originalA": left, "originalB": right,
                            "workerDecision": decision, "ownerCorrectionRequest": request,
                            "status": "원문 담당자 정정 요청(원문 변경 없음)"})

    def spec_original(rid):
        rec = specs[rid]
        return {"source": rec["source"], "text": f"{rid}: {rec['title']} / {rec['priority']} / {rec['status']}"}

    def definition_original(definition):
        fields = definition["fields"]
        return {"source": definition["source"], "text": fields.get("기능명", fields.get("요구사항명", "")),
                "rawAcceptance": fields.get("수용 기준", "")}

    def task_original(wid):
        task = tasks[wid]
        return {"source": task["source"], "text": f"{wid}: {task['title']} / {task['type']} / {task['start']}~{task['end']} / 가중치 {task['weight']} / raw {cell(task['rawSpecIds'])}"}

    for rid, rec in specs.items():
        for qid in rec["requirementIds"]:
            if len(requirements[qid]) == 1 and requirements[qid][0]["title"] != rec["title"]:
                requirement = requirements[qid][0]
                binding = req_bindings.get(rid)
                correction("REQUIREMENT_BINDING_BY_MEANING" if binding else "REQUIREMENT_TITLE_REVIEW", rid, spec_original(rid),
                           {"source": requirement["source"], "text": f"요구사항 {qid}: {requirement['title']}"},
                           f"의미 기준 요구 {binding['effectiveRequirementIds'] or '없음'}로 연결한다({binding['relation']}). 원문 번호는 보존한다." if binding
                           else "원문 요구 번호와 의미 연결을 보존하며 제목 차이만으로 자동 재연결하지 않는다.",
                           "명세/요구사항 담당자가 목차의 요구 번호를 현재 의미에 맞게 정정해 주세요." if binding
                           else "명세/요구사항 담당자가 분리·통합·번호 이동인지 확인하고 현재 ID와 의미별 연결을 정정해 주세요.")
        history, current = historical_definitions(rec), current_definitions(rec)
        if rid in active_ids and history and current:
            for definition in history:
                correction("HISTORICAL_DETAIL_SUPERSEDED", rid, definition_original(definition), definition_original(current[0]),
                           "옛 탭 정의는 이력으로 보존하고 현행 상세는 대체 시트로 사용한다. 번호 일치만으로 두 정의를 합치지 않는다.",
                           "명세 담당자가 옛 탭의 이력 표기와 `변경 내역` 시트를 현행 목차의 상세 시트에 맞춰 정정해 주세요.")
        elif len(current) > 1:
            for definition in current[1:]:
                correction("MULTIPLE_SPEC_DEFINITIONS", rid, definition_original(current[0]),
                           definition_original(definition), "두 상세를 독립 행으로 유지하고 자동으로 우선 정의를 선택하지 않는다.",
                           "명세 담당자가 두 정의의 유지/이전/통합 관계와 의미별 수용 기준을 명시해 주세요.")
        if rec.get("scopeClass") == "proposal":
            wid = decisions.get(rid, {}).get("wbsIds", [None])[0]
            if wid in tasks:
                correction("RAW_SCOPE_STATUS_CORRECTION", rid, spec_original(rid), task_original(wid),
                           f"유효 범위는 {effective_scope}이며 재사용 공수는 중복 산정하지 않는다.",
                           "명세/WBS 담당자가 미갱신 상태·우선순위 표기를 정정하고 기존 기능 재사용과 상세 단계의 일정·역할을 일치시켜 주세요.")
    for reserved in reserved_requirements:
        requirement = requirements[reserved["requirementId"]][0]
        correction("REQUIREMENT_MEANING_CONFLICT", f"요구{reserved['requirementId']} / {reserved['rawIndexSpecId']}", spec_original(reserved["rawIndexSpecId"]),
                   {"source": requirement["source"], "text": f"요구사항 {reserved['requirementId']}: {requirement['title']}"},
                   f"{reserved['effectiveStatus']}. {reserved['reason']}",
                   "명세/요구 담당자가 재무·실적의 고유 FS ID·목차·담당 WBS를 정하거나 제외를 공식 결정해 주세요.")
    for item in legacy_ai:
        if item["relation"].startswith("같은 ID"):
            continue
        legacy_rec = specs[item["legacySpecId"]]
        legacy_definition = next(d for d in legacy_rec["definitions"] if d["source"]["row"] == item["legacySource"]["row"] and definition_file(d) == item["legacySource"]["file"])
        correction("LEGACY_AI_ID_MEANING_MOVED", f"{item['legacySpecId']} → {', '.join(item['currentSpecIds'])}", definition_original(legacy_definition),
                   spec_original(item["currentSpecIds"][0]),
                   f"{item['relation']}. 보존 의무: {item['preserved']}",
                   "AI 명세 담당자가 같은 번호가 다른 의미로 재사용된 점을 목차의 FS051 상세 위치·요구 번호와 WBS FE 복합 산출물 열(BE 산출물·변경 내역은 10-12에 정정)에 명시하고 옛 FS046의 서버 공통 의무를 v2 상세에 반영할지 정해 주세요.")
    for rid, link in raw_corrections.items():
        correction("MALFORMED_SPEC_ID", f"{rid} → {link['effectiveSpecId']}", spec_original(rid), spec_original(link["effectiveSpecId"]),
                   "원문 기능 ID 셀은 고치지 않는다. 제목 행과 목차의 통합 선언에 근거해 FS066-01에만 유효 연결한다.",
                   f"명세 담당자가 {link['file']} {link['idRow']}행의 기능 ID 셀을 {link['effectiveSpecId']}로 정정해 주세요.")
    for note in links.get("wbsNotes", []):
        if all(wid in tasks for wid in note["wbsIds"]):
            correction(note["code"], " / ".join(note["wbsIds"]), task_original(note["wbsIds"][0]), task_original(note["wbsIds"][1]),
                       note["workerDecision"], note["ownerRequest"])
        else:
            finding("WBS_NOTE_BROKEN_ID", note["code"], "WBS 메모의 작업 번호가 현재 원문에 없다.", "error")
    for wid, task in tasks.items():
        if re.fullmatch(r"\d{4}-\d\d-\d\d", task["start"]) and re.fullmatch(r"\d{4}-\d\d-\d\d", task["end"]) and task["start"] > task["end"]:
            correction("WBS_DATE_ORDER_INVALID", wid, task_original(wid), {"source": task["source"], "text": "같은 행의 종료일이 시작일보다 앞선다."},
                       "날짜를 추정하거나 고치지 않고 일정 정정 대상으로 표시한다.", "WBS 일정 담당자가 시작일·종료일을 정정해 주세요.")
    if "6.5.3" in tasks and "FS035" in specs:
        correction("WBS_RAW_ID_REVIEW", "FS057 / 6.5.3", task_original("6.5.3"), spec_original("FS035"),
                   "기존 FS057→FS035 통합 선언에 따라 저장 업무를 FS035로 편집 연결하고 raw FS057은 보존한다.",
                   "WBS 담당자가 산출물 FS057을 현재 통합 ID와 일치시키고 기존 저장 업무와 중복 산정하지 않도록 정정해 주세요.")
    for records in catalog.get("wbsSources", {}).get("ancillaryRows", {}).values():
        for record in records:
            if any("6.5.13" in text and "6.5.20" in text for text in record["cells"]) and "6.6.13" in tasks:
                correction("WBS_COVER_BODY_CONFLICT", "FS074~FS081 / 6.6.13~20",
                           {"source": record["source"], "text": cell(record["cells"])}, task_original("6.6.13"),
                           "본문 6.6.13~20을 현재 역할 연결에 사용한다. 일정 기준은 WBS이며 내부 충돌은 정정 요청으로 남긴다.",
                           "WBS 담당자가 표지/요약/본문의 번호·역할·기간·가중치를 일치시켜 주세요. 원문 제안 8개를 제외하지 마세요.")
    correction_rows = [[r["code"], r["target"], source_label(r["originalA"]["source"]), r["originalA"]["text"],
                        source_label(r["originalB"]["source"]), r["originalB"]["text"],
                        r["workerDecision"], r["ownerCorrectionRequest"], r["status"]] for r in corrections]

    trace_rows, screen_rows, feature_rows, api_rows = [], [], [], []
    for rid, rec in specs.items():
        impls = implementation_records(rid)
        status = implementation_text(impls, "status", rid) or "문서 판정 없음"
        disposition = decisions.get(rid, {}).get("apiDisposition", "연결됨" if by_spec[rid] else "미연결")
        codes = list(dict.fromkeys([c for r in impls for c in r.get("codeLocations", [])] + (code_by_spec[rid] if rid in active_ids else [])))
        meaning_requirements = spec_mapping(rid, rec, detail_ctx, req_bindings)[1]
        matched = by_spec[rid] or [("", "", {})]
        for path, method, op in matched:
            x = op.get("x-zoop", {})
            trace_rows.append([rid, rec["title"], rec["priority"], rec["status"], rec["requirementIds"], meaning_requirements, x.get("screenIds", decisions.get(rid, {}).get("screenIds", [])),
                               op.get("operationId", ""), method.upper(), path, x.get("wbsIds", decisions.get(rid, {}).get("wbsIds", [])), status, disposition,
                               decisions.get(rid, {}).get("reason", ""), source_label(rec["source"]), implementation_text(impls, "source", rid), codes, rec.get("scopeClass", "existing"),
                               spec_scope(rid), authority.get("decisionId"), authority.get("source")])
        wids = list(dict.fromkeys(human_tasks[rid] + decisions.get(rid, {}).get("wbsIds", []) + (list(common_wbs) if by_spec[rid] else [])))
        for wid in wids or [""]:
            task = tasks.get(wid, {})
            feature_rows.append([rid, rec["title"], rec["priority"], wid, task.get("title", ""), task.get("start", ""), task.get("end", ""),
                                 task.get("owners", ""), status, disposition, implementation_text(impls, "detail", rid), codes, source_label(task.get("source")), implementation_text(impls, "source", rid),
                                 rec.get("scopeClass", "existing"), task.get("role", ""), task.get("rawSpecIds", []), wbs_bindings.get(wid, {}).get("reason", ""),
                                 spec_scope(rid), authority.get("decisionId")])
    for sid, screen in screens.items():
        for path, method, op in by_screen[sid] or [("", "", {})]:
            x = op.get("x-zoop", {})
            for rid in x.get("fsIds", []) + x.get("nfsIds", []) or [""]:
                screen_rows.append([sid, screen["title"], screen["route"], screen["page"], rid, specs.get(rid, {}).get("title", ""),
                                    op.get("operationId", ""), method.upper(), path, x.get("priority", ""), x.get("wbsIds", [])])
        for rid, decision in decisions.items():
            if sid in decision.get("screenIds", []):
                screen_rows.append([sid, screen["title"], screen["route"], screen["page"], rid, specs.get(rid, {}).get("title", ""),
                                    decision.get("apiDisposition", ""), "", "", specs.get(rid, {}).get("priority", ""), decision.get("wbsIds", [])])
    for path, method, op in all_ops:
        x = op["x-zoop"]
        api_rows.append([op["operationId"], op["summary"], method.upper(), path, x["authority"], x["priority"], x["fsIds"], x["nfsIds"],
                         x["screenIds"], x["wbsIds"], x["prototype"].get("path"), x["prototype"]["differences"], x["teamQuestions"],
                         x.get("requirementIds", []), x.get("featureStatus", {}).get("state", "")])
    proposals = [r for r in specs.values() if r.get("scopeClass") == "proposal"]
    proposal_rows = [[r["id"], r["title"], r["priority"], r["status"], decisions.get(r["id"], {}).get("reusesSpecIds", []),
                      decisions.get(r["id"], {}).get("wbsIds", []), decisions.get(r["id"], {}).get("apiDisposition", ""),
                      decisions.get(r["id"], {}).get("operationalIntegration", ""), effective_scope,
                      authority.get("decisionId"), authority.get("source"),
                      decisions.get(r["id"], {}).get("reason", ""), source_label(r["source"])] for r in proposals]
    def wbs_scope(wid):
        held = [rid for rid in wbs_bindings.get(wid, {}).get("specIds", []) if rid in exception_specs]
        return effective_scope + (f" (단 {', '.join(held)}의 사용자 기능 제외·API 유보 부분은 활성 범위 아님)" if held else "")

    wbs_rows = [[wid, t["title"], t.get("role", ""), t["type"], t["deliverable"], t["rawSpecIds"],
                 wbs_bindings.get(wid, {}).get("specIds", []), t["owners"], t["start"], t["end"], t["weight"], t["progress"],
                 wbs_scope(wid),
                 source_label(t["source"])] for wid, t in tasks.items()]
    authority_rows = []
    for e in exceptions:
        authority_rows.append(["범위 예외", e["id"], e["effectiveStatus"], e["basis"],
                               f"명세 {e.get('specIds') or '없음'} · 요구 {e.get('requirementIds', [])} · 유보 작업 {e.get('retainedOperationIds', [])}", e["reason"], authority.get("decisionId")])
    for r in reserved_requirements:
        requirement = requirements[r["requirementId"]][0]
        authority_rows.append(["유보 요구", f"요구 {r['requirementId']} {requirement['title']}", r["effectiveStatus"], source_label(requirement["source"]),
                               f"목차 {r['rawIndexSpecId']}(번호만 일치) · 관련 요구 {r['relatedRequirementIds']} · 작업 {r['operationIds']} · WBS {r['wbsIds'] or '없음'}", r["reason"], authority.get("decisionId")])
    for sheet in historical_files.values():
        authority_rows.append(["이력 시트", sheet["file"], f"이력(현행 대체 {sheet['supersededBy']})", sheet["sheet"], sheet["specIds"], sheet["reason"], authority.get("decisionId")])
    for rid in sorted(historical_only):
        mapped = sorted({t for m in legacy_ai if m["legacySpecId"] == rid for t in m["currentSpecIds"]})
        authority_rows.append(["이력 전용 ID", rid, spec_scope(rid), source_label(specs[rid]["source"]), mapped, "현재 목차에서 사라졌다. 활성 ID로 승격하지 않는다.", authority.get("decisionId")])
    for rid, link in raw_corrections.items():
        authority_rows.append(["원문 ID 정정 연결", rid, spec_scope(rid), source_label(specs[rid]["source"]), [link["effectiveSpecId"]], f"{link['observed']} {link['reported']} {link['status']}", authority.get("decisionId")])
    for link in detail_ctx["resolvedCorrections"]:
        authority_rows.append(["원문 ID 정정 해소", f"{link['rawId']} → {link['effectiveSpecId']}", f"원문 정정됨 · 활성 ID {link['effectiveSpecId']}",
                               f"{link['file']} {link['idRow']}행(제목 {link['headingRow']}행)", [link["effectiveSpecId"]],
                               f"{link['observed']} {link['boundary']} 이력 증거: {', '.join(link['historicalProof'])}", authority.get("decisionId")])
    for m in legacy_ai:
        authority_rows.append(["AI 의미 이행", f"{m['legacySpecId']} {m['legacyTitle']}", m["relation"], f"{m['legacySource']['file']}:{m['legacySource']['row']}", m["currentSpecIds"], f"요구 {m['requirementIds']} · 보존: {m['preserved']}", authority.get("decisionId")])
    for b in links.get("requirementBindings", []):
        authority_rows.append(["요구 의미 연결", b["specId"], b["relation"], f"목차 요구 {b['rawRequirementIds']}", b["effectiveRequirementIds"], b["note"], authority.get("decisionId")])
    warning_rows = [[w["code"], w["severity"], w["target"], w["message"], w["source"]] for w in findings]
    error_rows = [[r["code"], r["httpStatus"], r["meaning"], r["recovery"]] for r in api["x-zoop-error-codes"]]
    sheets = {
        "추적표": (["명세 ID", "명세명", "원본 우선순위", "원본 상태", "요구사항 번호(원문)", "의미 기준 요구사항", "화면 ID", "operationId", "메서드", "목표 경로", "WBS", "문서 구현 판정", "API 대응", "사유", "명세 출처", "판정 출처", "코드 위치", "원문 범위 분류", "유효 구현 범위", "유효 기준 결정", "유효 기준 출처"], trace_rows),
        "화면→기능→API": (["화면 ID", "화면명", "라우트", "캡처 쪽수", "명세 ID", "명세명", "operationId", "메서드", "목표 경로", "API 우선순위", "WBS"], screen_rows),
        "기능→WBS→구현 상태": (["명세 ID", "명세명", "원본 우선순위", "WBS", "업무", "시작일", "종료일", "담당", "문서 구현 판정", "API 대응", "판정 내용", "코드 위치", "WBS 출처", "판정 출처", "원문 범위 분류", "역할", "원문 FS/의존 참조", "편집 연결 근거", "유효 구현 범위", "유효 기준 결정"], feature_rows),
        "API 목록": (["operationId", "작업명", "메서드", "목표 경로", "권한", "API 우선순위", "FS", "NFS", "화면 ID", "WBS", "BFF 경로", "목표 차이", "팀 확인 필요", "요구사항 번호(FS 없음)", "기능 상태"], api_rows),
        "오류 코드": (["코드", "HTTP", "의미", "복구 규칙"], error_rows),
        "경고": (["분류", "수준", "대상 ID", "경고 내용", "출처"], warning_rows),
        "원문 WBS": (["WBS", "업무", "역할", "원문 유형", "원문 산출물", "rawSpecIds", "편집 명세 연결", "담당", "시작일", "종료일", "가중치", "진행률", "유효 구현 범위(원문과 별도)", "출처"], wbs_rows),
        "원문 제안·유효 기준": (["명세 ID", "명세명", "원본 우선순위", "원본 상태", "재사용 기존 FS", "현재 WBS", "API 대응", "운영 연동", "유효 구현 범위", "유효 기준 결정", "유효 기준 출처", "사유", "원문 출처"], proposal_rows),
        "범위 예외·상세 권한": (["구분", "대상", "유효 상태", "원문 출처", "현재 대응", "보존 내용·사유", "결정 출처"], authority_rows),
        "수용 기준": (["명세 ID", "상세 정의 번호", "목차명", "상세 기능명", "상세 정의 출처", "수용 기준 존재", "원문 수용 기준", "원본 우선순위", "원본 상태", "유효 구현 범위", "유효 기준 결정", "기존 문서 구현 판정", "판정 출처", "판정 적용 경계", "기존 증거 종류", "코드 위치(실행 증거 아님)", "현행 실행 검증", "상세 권한", "현행 대응 ID", "의미 기준 요구사항"], acceptance_rows),
    }
    readback = workbook(OUT / "zoop-trace.xlsx", sheets)
    print(f"xlsx 되읽기: {readback['status']} · {len(readback['sheetDataRows'])}시트 · {readback['cells']}셀")
    (OUT / "acceptance-matrix.json").write_text(json.dumps({"coverage": acceptance_coverage,
        "authorityDecisionId": authority.get("decisionId"), "records": acceptance_records}, ensure_ascii=False, indent=2) + "\n")
    (OUT / "acceptance-matrix.md").write_text("# 전체 명세의 상세 정의별 수용 기준\n\n원문 상세 정의를 각각 한 행으로 유지한다. 기존 ID 단위 구현 판정은 인용이며 해당 상세의 수용 기준 통과 증거가 아니다. **현행 실행 검증은 전부 미실행**이다. 상세/수용 기준이 없는 항목도 표시한다.\n\n"
        + f"- 범위: {len(specs)} ID(현행 목차 {len(active_ids)} + 이력 전용·원문 오류 {len(specs) - len(active_ids)}) · {actual_definitions}/{expected_definitions} 상세 정의(현행 {acceptance_coverage['currentDefinitions']} · 이력 {acceptance_coverage['historicalDefinitions']}) · {len(acceptance_records)}행\n"
        + f"- coverage: {acceptance_coverage['status']} · 수용 기준 없는 상세 {acceptance_coverage['missingAcceptanceDefinitions']}개 · 이력 정의는 현행 범위나 수용 기준 대체 증거가 아니다\n\n"
        + table(sheets["수용 기준"][0], acceptance_rows))
    (OUT / "source-corrections.json").write_text(json.dumps({"status": "원문 변경 없음",
        "authorityDecisionId": authority.get("decisionId"), "records": corrections}, ensure_ascii=False, indent=2) + "\n")
    (OUT / "source-corrections.md").write_text("# 원문 정정 요청\n\n두 원문·출처와 작업자 가결정을 분리한다. 이 요청으로 원문 ID/상태/일정을 수정하거나 기능 완료로 승격하지 않는다. WBS 내부 마일스톤/상위 종료일 충돌은 [10-08 기준](../../share/day-work-2026-10-08.md)의 일정 정정 요청과 [현재 작업 기준](../../share/day-work-2026-10-10.md)의 10-10 갱신을 함께 읽는다. [10-12 원문 갱신 기록](../../evidence/doc-delta-2026-10-12.md)은 10-10 이후 작업자 기록 변경과 팀 미기록 변경, 해소된 원문 ID 오류와 남은 정정 요청을 구분한다.\n\n"
        + table(["분류", "대상", "원문 A 출처", "원문 A", "원문 B 출처", "원문 B", "작업자 가결정", "담당자 수정 요청", "상태"], correction_rows))
    (OUT / "effective-authority.json").write_text(json.dumps({
        "authority": authority, "valid": authority_valid,
        "sourceSha256": hashlib.sha256(authority_source.read_bytes()).hexdigest() if authority_source.is_file() else None,
        "scopeExceptions": exceptions,
        "reservedRequirements": reserved_requirements,
        "detailAuthority": {"historicalSheets": list(historical_files.values()), "historicalOnlySpecIds": sorted(historical_only),
                            "rawIdCorrections": list(raw_corrections.values()), "resolvedRawIdCorrections": detail_ctx["resolvedCorrections"],
                            "activeIndexSpecIds": sorted(active_ids)},
        "specs": effective_specs,
        "wbsImplementationScope": {"status": effective_scope, "count": len(tasks)},
    }, ensure_ascii=False, indent=2) + "\n")
    (OUT / "detail-authority.json").write_text(json.dumps({"decisionId": authority.get("decisionId"), "rows": [dict(zip(sheets["범위 예외·상세 권한"][0], row)) for row in authority_rows]}, ensure_ascii=False, indent=2) + "\n")
    (OUT / "detail-authority.md").write_text("# 범위 예외 · 상세 시트 권한 · AI 의미 이행\n\n현행 상세와 이력 정의, 번호가 재사용된 AI ID, 요구 번호와 의미의 차이, 유보 범위를 한곳에 둔다. 원문 ID·상태·시트는 고치지 않으며 구현 완료·수용 기준 통과를 뜻하지 않는다.\n\n"
        + table(sheets["범위 예외·상세 권한"][0], authority_rows))
    (OUT / "warnings.md").write_text("# 추적 검사 결과\n\n원본 오류·미결정·목표/BFF 차이는 숨기지 않는다. 차이는 오류 수정 완료 또는 정식 구현 완료를 뜻하지 않는다.\n\n" + table(sheets["경고"][0], warning_rows))
    (OUT / "api-differences.md").write_text("# 목표 계약 ↔ BFF 관측본\n\n경로·요청/응답 필드·필수 여부·형식·권한을 비교한다. 각 작업의 전체 필드 차이는 `api-differences/<operationId>.json`에 있다. 실제 동작 시험이 아니며 현재 서버 조회 시각은 `bff-observation.json`에 있다.\n\n" + table(["operationId", "목표 경로", "BFF 경로", "필드 차이", "설계 차이"], differences))
    (OUT / "checks.json").write_text(json.dumps({"findings": findings, "counts": dict(Counter(w["code"] for w in findings)),
        "severityCounts": dict(Counter(w["severity"] for w in findings)),
        "xlsxReadback": readback,
        "acceptanceCoverage": acceptance_coverage,
        "sourceCorrectionCount": len(corrections),
        "authority": {"decisionId": authority.get("decisionId"), "source": authority.get("source"), "valid": authority_valid},
        "scopeCounts": {"rawIndexSpecs": len(specs) - len(historical_only) - len(raw_corrections),
                        "rawExistingSpecs": len(specs) - len(historical_only) - len(raw_corrections) - len(proposals), "rawProposals": len(proposals),
                        "rawDetailOnlyRecords": len(historical_only) + len(raw_corrections),
                        "historicalDetailOnlyIds": len(historical_only), "rawMalformedIds": len(raw_corrections),
                        "activeIndexSpecs": len(active_ids),
                        "scopeExceptionSpecs": len(exception_specs),
                        "effectiveIncludedSpecs": len(active_ids) - len(exception_specs) if authority_valid else 0,
                        "effectiveIncludedRawProposals": len(proposals) if authority_valid else 0,
                        "effectiveIncludedWbs": len(tasks) if authority_valid else 0,
                        "reservedRequirements": len(reserved_requirements),
                        "retainedOperationsOutsideActiveScope": sorted({oid for e in exceptions for oid in e.get("retainedOperationIds", [])}),
                        "wbs": len(tasks), "wbsRoles": dict(Counter(t.get("role", "") for t in tasks.values())),
                        "targetOperations": len(all_ops)},
        "sheetDataRows": {k: len(v[1]) for k, v in sheets.items()}}, ensure_ascii=False, indent=2) + "\n")
    migration = links.get("wbsMigration", {})
    (OUT / "wbs-migration.md").write_text("# WBS 내용 기반 연결 이행\n\n기계 추출 raw와 아래 편집 관계는 별개다. 번호 존재만으로 옛 의미를 승계하지 않는다. 승인·구현 상태를 자동 변경하지 않는다.\n\n"
        + f"- 이전: `{migration.get('previousWbsSource', '')}` · {migration.get('previousTaskCount', '')}개\n"
        + f"- 현재: `{migration.get('currentWbsSource', '')}` · {len(tasks)}개 · 역할 {dict(Counter(t.get('role', '') for t in tasks.values()))}\n"
        + f"- 헤더 대조: {migration.get('headerComparison', '')}\n"
        + f"- 목표 API {len(all_ops)}개를 현재 업무명·역할·FS 편집 연결로 검사했다. 작업별 연결 정본은 openapi.yaml의 x-zoop.wbsIds다.\n"
        + f"- 원문 제안 {len(proposals)}개도 현재 유효 기준 '{effective_scope}'에 포함한다. 결정 출처는 `{authority.get('source', '')}`다. 표지/요약 vs 본문 및 원문 ID 차이는 정정 요청으로 warnings.md에 유지한다. 재사용 공수를 중복 산정하지 않는다.\n\n"
        + table(["이전 번호", "의미", "현재 역할/번호", "내용 대조 근거"],
                [[g["previousIds"], g["meaning"], g["currentTargets"], g["reason"]] for g in migration.get("groups", [])]))
    for sid, screen in screens.items():
        lines = f"# {sid} · {screen['title']}\n\n[전체 목록](index.md) · [검사 결과](warnings.md)\n\n- 라우트: `{screen['route'] or '설명 슬라이드(라우트 없음)'}`\n- 화면계획서 캡처 쪽수: **{screen['page']}**\n- 원본: {source_link(screen['source'], 3)}\n\n"
        rows_out = []
        for path, method, op in by_screen[sid]:
            fs_links = " · ".join(f"[{rid}](feature-{slug(rid)}.md)" for rid in op["x-zoop"]["fsIds"])
            rows_out.append([fs_links, f"[{op['operationId']}](../../api/out/zoop-api-spec.md#{op['operationId'].lower()})", f"{method.upper()} {path}", op["x-zoop"]["priority"]])
        lines += table(["기능", "목표 API", "경로", "우선순위"], rows_out) if rows_out else "이 쪽은 화면 규칙/설명 슬라이드이거나 현재 API 연결이 없는 화면이다. API를 임의로 만들지 않는다.\n"
        quality = list(dict.fromkeys(rid for _, _, op in by_screen[sid] for rid in op["x-zoop"]["nfsIds"]))
        if quality:
            lines += "\n공통 품질: " + " · ".join(f"[{rid}](feature-{slug(rid)}.md)" for rid in quality) + "\n"
        screen_only = [rid for rid, decision in decisions.items() if sid in decision.get("screenIds", [])]
        if screen_only:
            lines += "\n화면 구현 요구: " + " · ".join(f"[{rid}](feature-{slug(rid)}.md)" for rid in screen_only) + "\n"
        (OUT / f"screen-{slug(sid)}.md").write_text(lines)
    for rid, rec in specs.items():
        impls = implementation_records(rid)
        decision = decisions.get(rid, {})
        feature_screens = list(dict.fromkeys([sid for _, _, op in by_spec[rid] for sid in op["x-zoop"]["screenIds"]] + decision.get("screenIds", [])))
        meaning_requirements = spec_mapping(rid, rec, detail_ctx, req_bindings)[1]
        lines = f"# {rid} · {rec['title']}\n\n[전체 목록](index.md) · [검사 결과](warnings.md) · [범위 예외·상세 권한](detail-authority.md)\n\n" + table(["항목", "원본/연결"], [["원본 우선순위·상태", f"{rec['priority']} · {rec['status']}"], ["원문 범위 분류", rec.get("scopeClass", "existing")], ["유효 구현 범위", spec_scope(rid)], ["유효 기준", f"{authority.get('decisionId')} · {authority.get('source')}"], ["완료·실행 경계", "요구 범위 포함은 구현 완료·유료 실행 승인·본개발 채택이 아님"], ["요구사항 번호(원문)", rec["requirementIds"]], ["의미 기준 요구사항", meaning_requirements], ["API 대응", decision.get("apiDisposition", "연결됨" if by_spec[rid] else "미연결")], ["사유", decision.get("reason", "")], ["프로토타입 문서 판정", implementation_text(impls, "status", rid) or "판정 없음"], ["코드 위치(실행 증거 아님)", list(dict.fromkeys([c for r in impls for c in r.get("codeLocations", [])] + (code_by_spec[rid] if rid in active_ids else [])))]])
        lines += "\n## 원본과 의미\n\n" + source_link(rec["source"], 3) + "\n\n"
        for definition in rec["definitions"]:
            fields = definition["fields"]
            description = fields.get("기능 설명", fields.get("요구사항 설명", ""))
            label = "[이력 정의] " if definition_file(definition) in historical_files else ""
            # 한 쪽 요약에는 원문 요약문만 발췌하고 상세 정책/수용 기준은 고정본으로 연결한다.
            lines += f"- {label}{source_link(definition['source'], 3)} — {fields.get('기능명', fields.get('요구사항명', ''))}: {description}\n"
        notes = []
        if rid in historical_only:
            notes.append("현재 목차 밖의 이력 ID다. 활성 목록으로 승격하지 않는다.")
        if rid in raw_corrections:
            notes.append(f"원문 기능 ID 셀은 정정되지 않았다. 제목 행에 근거해 {raw_corrections[rid]['effectiveSpecId']}에 유효 연결한다. {raw_corrections[rid]['status']}")
        for raw_id, link in raw_corrections.items():
            if link["effectiveSpecId"] == rid:
                notes.append(f"정정 연결 정의: 원문 ID `{raw_id}`의 {link['file']} {link['idRow']}행(제목 {link['headingRow']}행).")
        for item in legacy_ai:
            if item["legacySpecId"] == rid:
                notes.append(f"이 번호의 옛 AI 탭 정의는 현재 {', '.join(item['currentSpecIds'])}로 이행된다({item['relation']}). 보존: {item['preserved']}")
            elif rid in item["currentSpecIds"]:
                notes.append(f"옛 {item['legacySpecId']} {item['legacyTitle']}에서 이행됨({item['relation']}). 보존: {item['preserved']}")
        if rid in exception_specs:
            notes.append(f"{exception_specs[rid]['effectiveStatus']}: {exception_specs[rid]['reason']}")
        for reserved in reserved_requirements:
            if reserved["rawIndexSpecId"] == rid:
                notes.append(f"목차의 요구 번호 {reserved['requirementId']}는 번호만 일치한다. 요구 {reserved['requirementId']}는 {reserved['effectiveStatus']}이며 {reserved['operationIds']}에 요구 번호로만 연결한다.")
        if notes:
            lines += "\n## 상세 권한과 의미 이행\n\n" + "".join(f"- {note}\n" for note in notes)
        if impls:
            lines += "\n문서 판정 출처: " + " / ".join(f"옛 {r.get('legacySpecId', rid)} {source_link(r['source'], 3)}" if r.get("legacySpecId", rid) != rid else source_link(r["source"], 3) for r in impls) + ". 기준선에 반영 현황의 마지막 행을 적용한 판정이며 번호가 재사용된 AI ID는 의미 이행표가 정한 곳에만 적용한다. 이번 생성기의 실행 검증은 아니다.\n"
        if decision.get("reusesSpecIds"):
            lines += "\n재사용 근거(상세 수용 기준의 구현 완료 증거 아님): " + " · ".join(f"[{target}](feature-{slug(target)}.md)" for target in decision["reusesSpecIds"]) + "\n"
        lines += "\n## WBS 편집 연결\n\n" + table(["WBS", "역할", "현재 업무", "연결 근거"],
            [[wid, tasks[wid].get("role", ""), tasks[wid]["title"], wbs_bindings.get(wid, {}).get("reason", decision.get("reason", ""))]
             for wid in dict.fromkeys(human_tasks[rid] + decision.get("wbsIds", [])) if wid in tasks])
        lines += "\n## 화면·목표 API\n\n" + " · ".join(f"[{sid}](screen-{slug(sid)}.md)" for sid in feature_screens) + "\n\n"
        if len(by_spec[rid]) <= 12:
            lines += table(["작업", "메서드·목표 경로", "WBS", "권한"], [[f"[{op['operationId']}](../../api/out/zoop-api-spec.md#{op['operationId'].lower()})", f"{method.upper()} {path}", op["x-zoop"]["wbsIds"], op["x-zoop"]["authority"]] for path, method, op in by_spec[rid]])
        else:
            groups = defaultdict(list)
            for path, method, op in by_spec[rid]:
                groups[path.split("/")[3]].append(op)
            lines += table(["도메인", "연결 작업 수", "작업 목록"], [[group, len(ops), "[전체 API 목록](../../api/out/zoop-api-spec.md#작업-목록)"] for group, ops in groups.items()])
            lines += "\n작업별 FS/NFS 역링크와 전체 연결은 API 명세서 및 xlsx 추적표에 보존한다.\n"
        feature_findings = [w for w in findings if w["target"] == rid]
        if feature_findings:
            lines += "\n## 확인 필요\n\n" + "".join(f"- {w['code']}: {w['message']}\n" for w in feature_findings)
        (OUT / f"feature-{slug(rid)}.md").write_text(lines)
    (OUT / "index.md").write_text("# 줍 연결 문서\n\n이 문서는 생성물이다. 변경은 catalog.yaml 가져오기, links.yaml 또는 docs/api/openapi.yaml 편집으로 한다.\n\n[현재 사용자 결정(10-10)](../../share/day-work-2026-10-10.md) · [이전 결정(10-08)](../../share/day-work-2026-10-08.md) · [10-12 원문 갱신 기록](../../evidence/doc-delta-2026-10-12.md) · [10-10 원문 갱신 근거](../../evidence/doc-delta-2026-10-10.md) · [API v2 검토·채택/정정 목록](../../evidence/api-v2-review-2026-10-08.md) · [유효 기준 JSON](effective-authority.json) · [범위 예외·상세 권한](detail-authority.md) · [전체 수용 기준](acceptance-matrix.md) · [원문 정정 요청](source-corrections.md) · [API 명세서](../../api/out/zoop-api-spec.md) · [검사 결과](warnings.md) · [WBS 내용 기반 이행](wbs-migration.md) · [BFF 차이](api-differences.md) · [팀용 xlsx](zoop-trace.xlsx)\n\n"
        + f"현행 목차 {len(active_ids)}개 ID의 유효 구현 범위: **{effective_scope}**(기록된 예외 {len(exception_specs)}개 명세 제외). 원문 상태는 아래에 별도로 보존하며 완료·실행 승인으로 해석하지 않는다.\n\n## 화면\n\n"
        + "\n".join(f"- [{sid} · {s['title']}](screen-{slug(sid)}.md) — {s['page']}쪽" for sid, s in screens.items())
        + "\n\n## 현행 목차의 기능·비기능(원문 추가 제안 제외)\n\n"
        + "\n".join(f"- [{rid} · {s['title']}](feature-{slug(rid)}.md) — 원문 우선순위 {s['priority']} · 유효 {spec_scope(rid)}" for rid, s in specs.items() if rid in active_ids and s.get("scopeClass") != "proposal")
        + "\n\n## 원문 추가 제안 8개 · 현재 팀합의 설계대상\n\n"
        + "\n".join(f"- [{r['id']} · {r['title']}](feature-{slug(r['id'])}.md) — 원문 {r['priority']} · {r['status']} / 유효 {effective_scope}" for r in proposals)
        + "\n\n## 현행 목차 밖 이력·원문 오류 ID(활성 목록 아님)\n\n"
        + "\n".join(f"- [{rid} · {s['title']}](feature-{slug(rid)}.md) — {spec_scope(rid)}" for rid, s in specs.items() if rid not in active_ids) + "\n")
    api_md = "# 줍 정식 서버 목표 API 명세서\n\n생성물이다. 정본은 [openapi.yaml](../openapi.yaml)이다. WBS **6.1.2 API 명세 최종 검토·구현 기준 확정**, **6.11.1 API·스키마 계약 통합 점검**에 연결한다. API 공통 규칙은 사용자 **잠정 채택/시연 기준**이며 정식 구현 완료가 아니다. 전체 명세/WBS(FS074~081 포함)의 현재 범위는 팀합의요구/설계대상이며 원문 승인 표기·구현 증거·유료 실행 승인은 별개다. 추천 종목·산업·테마(FS011-02)는 활성 사용자 기능에서 제외하되 작업을 유보하며 요구 42의 재무·실적은 폐기되지 않았다(기록된 예외). 프로토타입 관측과 구현 목표를 구분한다.\n\n[현재 사용자 결정(10-10)](../../share/day-work-2026-10-10.md) · [이전 결정(10-08)](../../share/day-work-2026-10-08.md) · [API v2 검토·채택/정정 목록](../../evidence/api-v2-review-2026-10-08.md) · [10-12 원문 갱신 기록](../../evidence/doc-delta-2026-10-12.md) · [10-10 원문 갱신 근거](../../evidence/doc-delta-2026-10-10.md) · [유효 기준 JSON](../../system/out/effective-authority.json) · [범위 예외·상세 권한](../../system/out/detail-authority.md) · [추적표](../../system/out/index.md) · [현재 검사 결과](../../system/out/warnings.md) · [WBS 내용 기반 이행](../../system/out/wbs-migration.md) · [BFF 차이](../../system/out/api-differences.md)\n\n## 공통 규칙\n\n"
    for name, rule in api["x-zoop-common-rules"].items():
        api_md += f"### {name}\n\n{rule}\n\n"
    api_md += "## 오류 코드\n\n" + table(sheets["오류 코드"][0], error_rows) + "\n## 작업 목록\n\n" + table(["작업", "메서드", "목표 경로", "권한", "우선순위"], [[f"[{op['operationId']}](#{op['operationId'].lower()})", method.upper(), path, op["x-zoop"]["authority"], op["x-zoop"]["priority"]] for path, method, op in all_ops])
    for path, method, op in all_ops:
        x = op["x-zoop"]
        fslinks = " · ".join(f"[{rid}](../../system/out/feature-{slug(rid)}.md)" for rid in x["fsIds"] + x["nfsIds"])
        slinks = " · ".join(f"[{sid}](../../system/out/screen-{slug(sid)}.md)" for sid in x["screenIds"])
        body = deref(api, op.get("requestBody", {}))
        body_schemas = [v.get("schema", {}) for v in body.get("content", {}).values()]
        response_schemas = []
        for status, response in op["responses"].items():
            resp = deref(api, response)
            response_schemas += [f"{status}: {schema_md(v.get('schema', {}))}" for v in resp.get("content", {}).values()]
            if not resp.get("content"):
                response_schemas.append(f"{status}: 본문 없음")
        api_md += f"\n<a id=\"{op['operationId'].lower()}\"></a>\n\n## {op['operationId']} · {op['summary']}\n\n{op['description']}\n\n" + table(["항목", "계약"], [["메서드·경로", f"`{method.upper()} {path}`"], ["권한·우선순위", f"{x['authority']} · {x['priority']}"], ["요청 본문", [schema_md(v) for v in body_schemas] or "없음"], ["응답", response_schemas], ["오류 코드", x["errorCodes"]], ["FS/NFS", fslinks], ["화면", slinks], ["WBS", x["wbsIds"]], ["프로토타입 대응", x["prototype"].get("path") or "없음"], ["목표와 차이", x["prototype"]["differences"]], ["팀 확인 필요", x["teamQuestions"] or "해당 작업의 별도 항목 없음. 공통 미결정은 별도로 남는다."]])
        if x.get("meaningQualifications"):
            api_md += "\n명세 의미 구분: " + " / ".join(f"{rid}: {meaning}" for rid, meaning in x["meaningQualifications"].items()) + "\n"
        if x.get("requirementIds"):
            api_md += "\n요구사항 번호 연결(FS 없음): " + " / ".join(f"{qid}: {x.get('requirementQualifications', {}).get(qid, '')}" for qid in x["requirementIds"]) + "\n"
        if x.get("featureStatus"):
            status = x["featureStatus"]
            api_md += f"\n기능 상태: `{status['state']}` · {status['decisionId']} · {status['note']}\n"
        parameter_rows = []
        for parameter in op.get("parameters", []):
            p = deref(api, parameter)
            parameter_rows.append([p["name"], p["in"], "필수" if p.get("required") else "선택", schema_label(p.get("schema", {})), p.get("description", "")])
        api_md += "\n### 요청 필드\n\n" + table(["이름", "위치", "필수", "형식", "설명"], parameter_rows)
        for schema in body_schemas:
            label = schema_label(schema)
            api_md += f"\n본문 스키마: [{label}](#schema-{slug(label).lower()}). 중첩 필드와 nullable/enum은 아래 공통 스키마를 확인한다.\n"
    api_md += "\n## 공통 스키마\n\nnullable은 미제공/미판정이고 숫자 0 또는 빈 문자열로 대신하지 않는다. 필수 필드는 null 허용 여부와 별개다. JSON 필드 이름·중첩 참조의 정본은 OpenAPI다.\n"
    for name, schema in api["components"]["schemas"].items():
        api_md += f"\n<a id=\"schema-{slug(name).lower()}\"></a>\n\n### {name}\n\n{schema.get('description', '')}\n\n"
        if schema.get("properties"):
            api_md += table(["필드", "필수", "형식/참조", "설명"], inline_schema_rows(schema))
        else:
            api_md += schema_md(schema) + "\n"
    (API_OUT / "zoop-api-spec.md").write_text(api_md)
    print(f"생성: 명세 {len(specs)}, 목표 API {len(all_ops)}, 화면/설명 슬라이드 {len(screens)}")
    print("xlsx 데이터 행: " + ", ".join(f"{name} {len(rows)}" for name, (_, rows) in sheets.items()))
    counts = Counter(w["code"] for w in findings)
    print("검사 목록: " + ", ".join(f"{name} {count}" for name, count in sorted(counts.items())))
    if args.check:
        for w in findings:
            print(f"[{w['severity']}] {w['code']} {w['target']}: {w['message']}")
        print(f"검사 종료: {len(findings)}건. 전체 필드 차이·출처는 docs/system/out/에 보존했다.")
        return 1 if findings else 0
    return 0


if __name__ == "__main__":
    sys.exit(main())
