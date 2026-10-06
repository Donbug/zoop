"""캡처 manifest로 16:9 PPTX를 만들고 LibreOffice PDF·페이지 PNG를 출력한다."""
from pathlib import Path
import json, subprocess, os, argparse, math, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
CAP=BASE/'captures'
TOK=json.loads((ROOT/'docs/design/tokens.json').read_text())['color']
C=lambda key: RGBColor.from_string(TOK['neutral'][key]['$value'].lstrip('#'))
BRAND=RGBColor.from_string(TOK['brand']['primary']['$value'].lstrip('#'))
FONT='SUIT'
prs=Presentation(); prs.slide_width=Inches(16); prs.slide_height=Inches(9)
index=[]


def box(slide,x,y,w,h,text,size=18,bold=False,color=None,fill=None,spacing=6,line_spacing=1.12):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb=fill
    tf=sh.text_frame; tf.word_wrap=True
    tf.margin_left=Inches(.03); tf.margin_right=Inches(.03); tf.margin_top=Inches(.03); tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.font.name=FONT; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color or C('text-primary')
        p.space_after=Pt(spacing)
        p.line_spacing=line_spacing
        for run in p.runs:
            rpr=run._r.get_or_add_rPr()
            ea=OxmlElement('a:ea'); ea.set('typeface',FONT); rpr.append(ea)
    return sh


def new(sid,title,route='',source=''):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=C('surface-base')
    box(slide,.6,.32,14.8,.52,f'{sid}   {title}',27,True)
    box(slide,.63,.92,14.8,.35,route,11,color=C('text-secondary'))
    line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(.6),Inches(1.27),Inches(15.4),Inches(1.27))
    line.line.color.rgb=C('border-subtle'); line.line.width=Pt(.8)
    box(slide,.63,8.45,13.7,.3,'줍 화면계획서 v2.1 · 프로토타입 기준'+('  |  '+source if source else ''),9,color=C('text-tertiary'))
    box(slide,14.6,8.4,.8,.4,f'{len(prs.slides):02}',12,bold=True,color=BRAND)
    index.append({'page':len(prs.slides),'id':sid,'title':title,'route':route})
    return slide


def columns(sid,title,lefttitle,left,righttitle,right,source=''):
    s=new(sid,title,source=source)
    box(s,.7,1.6,6.7,.45,lefttitle,21,True,color=BRAND)
    box(s,.7,2.25,6.7,5.65,left,18)
    box(s,8.3,1.6,6.8,.45,righttitle,21,True,color=BRAND)
    box(s,8.3,2.25,6.8,5.65,right,18)
    return s


def table(sid,title,headers,rows,widths,source=''):
    s=new(sid,title,source=source)
    x=.65; y=1.65
    for header,w in zip(headers,widths):
        box(s,x,y,w,.48,header,15,True,color=BRAND); x+=w
    y+=.62
    height=min(.77,5.65/max(1,len(rows)))
    for row in rows:
        x=.65
        for val,w in zip(row,widths):
            box(s,x,y,w-.15,height,str(val),13)
            x+=w
        line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(.65),Inches(y+height-.02),Inches(15.3),Inches(y+height-.02))
        line.line.color.rgb=C('border-subtle'); line.line.width=Pt(.6)
        y+=height
    return s


s=new('COVER','줍 화면계획서 v2.1','프로토타입 기준 · 2026-10-06')
box(s,.8,2.0,13.7,1.55,'실제 화면에서\n근거 확인까지',40,True,color=BRAND)
box(s,.8,4.3,13.5,2.7,'기능명세·요구사항 10-06 고정본 + 현재 프로토타입\n종목 메뉴 · 기업 전환 · 시점 관계 그래프 · 시간 버킷\n모의 계좌 + 실제 기사 저장본을 사용한 시연이에요.\n긍부정은 디자인시스템 예시만 있으며 데모 판정값은 없어요.',21)
columns('HISTORY','변경 이력·v1.7 대비','채택 6','1. 장 마감 보드의 회차 비교 → 퇴근 브리핑\n2. 모바일 아침 이후 바뀐 내용 → 신규·이어짐·빠짐\n3. 근거·원문 접근 상태 → 제공 불가와 실패 분리\n4. 차트·사건 시점 겹침 → 인과관계 아님 고지\n5. 모바일 AI 탭 → 기존 질문 창 열기\n6. 비회원 권한표·관리자 감사 규칙의 형식','불채택 6','1. D-GRP-06 간선 모달 → FS022·D03과 충돌\n2. 외부 원인 10종·공급/경쟁 직결선 → 데이터 없음\n3. 긍정/부정 가능 배지 → 방향 값 생산자 없음\n4. 모듈별 6색·카드 좌측 막대 → 색은 뜻에만\n5. 관심 최대 3개 → 무제한·0개 허용\n6. 08:00/18:00 → 07:40/18:10 생성 시작','작성 명세 §팀 v1.7 판정 · review.md 머리 Main 정정')
columns('HISTORY-21','변경 이력 · v2.0 → v2.1','화면·동작','1. 지도 목록을 독립 종목 메뉴로 분리\n2. 데스크톱 4메뉴·모바일 하단 5탭\n3. 기업 패널의 시세·추이·주요 지표\n4. 기업 상세의 좌측 전환·관계 그래프\n5. 시간 버킷·차트·그래프 선택 연동\n6. 차트 시점 선택 → asOf 공개 기준 변경','근거·문서','요구사항 10-06의 85개 레코드 번호로 갱신\nFS037 종목 메뉴를 데스크톱·모바일에 추가\n모바일 기업 전환 시트를 별도 화면으로 기록\n긍부정·회차 판단은 디자인시스템 예시 명시\n데모에는 실제 모델 호출·판정값이 없어요.\n부록 문서 체계와 기능·API·WBS 추적 연결','doc-delta-2026-10-06.md §2.2 · design-direction.md §5.3·§5.10')
columns('GUIDE','읽는 법','화면·번호','D: 데스크톱 · M: 모바일 · A: 관리자\nAUT 시작 / ONB 설정 / WAT 관심\nGRP 지도 / STK 종목 / COM 기업\nBRF 브리핑 / AI 질문 / MY 내 종목\nADM 운영 / SEN 디자인시스템 예시\nb/c: 추가 상태 · T2~T5: 튜토리얼 단계\n상자: 실제 조작 요소 · 원: 32px 번호\n한 화면의 주요 조작은 3~6개예요.','상태·근거','상태 칸: 그 화면에 구현된 상태\n기존 상태 캡처는 docs/design/ 파일로 참조\n코드 경로는 prototype/web/src/ 기준\n번호 행동: 실제 클릭 결과 또는 확인한 이동 참조\n모바일: 누르기 전 / 누른 뒤\n빈 상태와 요청 실패·권한 제한을 구분해요.\n실행에서 못 본 결과는 성공으로 쓰지 않아요.','captures/*.json clickEvidence · spec-v1.2-gap.md §0.3')

manifest=json.loads((CAP/'manifest.json').read_text())
SCREEN_ORDER=['D-AUT-01','D-AUT-02','D-AUT-03','D-ONB-01','D-ONB-02','D-WAT-01','D-WAT-01b',
    'D-GRP-01','D-GRP-01b','D-GRP-T','D-GRP-T2','D-GRP-T3','D-GRP-T4','D-GRP-T5',
    'D-GRP-02','D-GRP-03','D-GRP-04','D-GRP-05','D-GRP-06','D-GRP-07','D-GRP-07b',
    'D-STK-01','D-COM-01','D-COM-01b','D-COM-02','D-COM-03','D-COM-04','D-COM-05','D-COM-06','D-MY-01','D-BRF-01','D-BRF-02','D-BRF-04',
    'D-AI-01','D-AI-01b','D-AI-02','D-AI-02b','M-GRP-01','M-STK-01','M-BRF-01','M-AI-01','M-COM-01','M-COM-02','D-SEN-01','D-SEN-02',
    'A-ADM-05c','A-ADM-01','A-ADM-02','A-ADM-02b','A-ADM-03','A-ADM-04','A-ADM-04b','A-ADM-05','A-ADM-05b']
assert {x['id'] for x in manifest}==set(SCREEN_ORDER)
manifest.sort(key=lambda x:SCREEN_ORDER.index(x['id']))
ui={'AUT':'UI-01','ONB':'UI-02','WAT':'UI-02','GRP':'UI-03','STK':'FS037','COM':'UI-06/07','BRF':'UI-08','AI':'UI-10','MY':'신규','SEN':'예시 전용'}
ui_detail={'D-AUT-03':'개발용','D-GRP-02':'UI-04','D-GRP-03':'UI-04','D-GRP-04':'UI-05','D-GRP-05':'UI-09','D-GRP-06':'UI-03+04','M-GRP-01':'UI-03+04','D-BRF-04':'UI-03'}
rows=[]
for item in manifest:
    ident=item['id']; mod=ident.split('-')[1]
    route=item['route']
    # 주소는 실제 캡처 주소를 유지하되 길이가 긴 대상 id는 일반화한다.
    if mod=='GRP': route='/?mode=…'+ ('&impact=…' if 'impact=' in route else '&path=…' if 'path=' in route else '&round=…' if 'round=' in route else '&select=…' if 'select=' in route else '&tour=1' if 'tour=' in route else '')
    if mod=='COM': route='/companies/:id?period=…&asOf=…&tab=…'+('&issue=…' if 'issue=' in item['route'] else '')
    if mod=='BRF': route='/briefing?date=…&edition=…' if ident!='D-BRF-04' else '/?round=…'
    if mod=='AI': route='/?select=… (AI 질문 창)' if 'select=' in item['route'] else '/ (AI 질문 창)'
    rows.append((ident,route,ui_detail.get(ident,ui.get(mod,'관리자 전용')),' · '.join(item['fs'].split('·'))))
for n in range(0,len(rows),10):
    table(f'MAP-{n//10+1:02}','화면 목록·라우트·명세',['화면 ID','라우트(실제 진입 표면)','기존 문서 ID','명세'],rows[n:n+10],[2.3,5.2,2.6,4.55],'App.tsx · design-direction.md §5·§6.3 · spec-v1.2-gap.md')

s=new('FLOW','첫 진입 → 첫 근거 확인','현재 그래프 기본 보기 기준. 과거 문서의 2클릭 수치를 그대로 승계하지 않아요.','design-direction.md §5.2·§6.1·§6.4')
def node(x,y,w,text):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(.85))
    sh.fill.solid(); sh.fill.fore_color.rgb=C('surface-raised'); sh.line.color.rgb=C('border-strong')
    box(s,x+.12,y+.12,w-.24,.58,text,16,True)
def arrow(x,y,x2,y2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x2),Inches(y2)); sh.line.color.rgb=BRAND; sh.line.width=Pt(2)
    end=OxmlElement('a:tailEnd'); end.set('type','triangle'); sh.line._get_or_add_ln().append(end)
node(.8,1.8,2.8,'시작 /login'); node(4.4,1.8,3.4,'비회원 체험 → 시장 전체'); node(4.4,3.25,3.4,'회원 → 설정 완료 확인'); node(8.6,3.25,3.4,'미완료 → 투자·계좌·관심'); node(8.6,1.8,3.4,'기업 또는 사건 선택'); node(8.6,5,3.4,'실제 지도 튜토리얼 5단계'); node(12.6,1.8,2.5,'근거 창 확인')
arrow(3.6,2.2,4.4,2.2); arrow(2.2,2.65,4.4,3.7); arrow(7.8,2.2,8.6,2.2); arrow(7.8,3.7,8.6,3.7); arrow(10.3,4.1,10.3,5); arrow(12,2.2,12.6,2.2)
arrow(7.5,3.25,8.6,2.5); box(s,7.55,2.83,1,.28,'완료',11,color=BRAND)
arrow(12,5.4,12.25,5.4); arrow(12.25,5.4,12.25,2.65); arrow(12.25,2.65,11.85,2.65)
box(s,.9,6.75,14.2,1.05,'기업 패널의 직접 이슈 → 영향 경로 → 근거 · 사건 패널 → 근거\n근거 창 닫기 → 원래 선택 유지 · 기업 상세의 그래프로 복귀 → 지도 상태 복원',19)
columns('SEMANTICS','색과 아이콘의 뜻','색은 의미를 구분해요','브랜드: 주 행동·선택·포커스\n중성 회색: 내용·구분선·보조 정보\n상승/하락: 부호·▲/▼·수치와 함께\n경고/오류: 문구·아이콘과 함께\n모듈은 색이 아니라 ID·아이콘으로 구분해요.\n캡처 안 기업 상징색은 기업 구별용이에요.','같은 뜻에 같은 아이콘','Map 지도 · Globe 시장 전체\nBookmark 내 종목 · Star 관심 · Wallet 보유\nBuilding2 기업 · Zap 사건\nGitFork 간접 이슈 · Link2 연결\nFileText 근거 · Sparkles AI · Info 설명\n색만으로 상태나 관계를 전달하지 않아요.','tokens.json v3.0.0 · design-direction.md §4.6')
columns('STATE-01','공통 상태 규칙','조회·자료','로딩: 조회 중 표시, 다른 영역은 유지\n빈: 자료 없음 + 가능한 다음 행동\n부분 실패: 받은 자료 유지, 실패만 재조회\n오류: 사유와 해당 조회의 다시 시도\n자료 삭제·비공개·권한 없음: 제공 불가\n404·403 근거를 일시 장애로 재시도하지 않아요.','권한·작업','비회원: 시장 전체·근거 확인 가능\n내 종목·AI 전송: 로그인 안내\nAI: 결과 회수와 새 작업 실행은 별개\n발행 브리핑: 지표 한 개 실패도 본문은 유지\n지표 미제공과 실패를 글자로 구분해요.\n이 슬라이드는 소스 규칙이며 전 상태 캡처는 아니에요.','spec-v1.2-gap.md §0.2·§0.3 · design-direction.md §5.5')
table('M-RULE-01','모바일 전환 규칙',['PC 요소','모바일 표면','동작'],[
('헤더 메뉴','하단 5탭','지도 · 종목 · 브리핑 · AI · 내 종목'),('기업·사건 오른쪽 패널','바텀시트','절반/확장, 닫기 제공'),('모드 옆 범위 필터','모드 아래 전체 폭','전체 · 보유 · 관심'),('지도 도구 가로 줄','오른쪽 아래 세로 버튼','탭바 위 44px 조작 영역'),('시장 지표 3열','시장 지표 3행','관측 시각·출처·상태 유지'),('기업 상세 좌측 전환','기업 전환 시트','검색·보유·관심·최근·연결 기업'),('기업 상세 오른쪽 관계','본문 아래 관계 그래프','같은 기준 시점·선택 사건 유지'),('AI에게 묻기','AI 하단 탭','화면 이동 없이 질문 창 열기')],[4.1,4.3,6.25],'design-direction.md §4.8·§5.3·§5.10 · M-* 캡처')
table('AUTHORITY','비회원·회원·관리자 권한',['행동','비회원','회원','관리자 세션'],[
('시장 전체·기업·근거 조회','가능','가능','회원 세션과 분리'),('내 종목 지도·관심 저장','로그인 안내','가능','회원 세션과 분리'),('모의 계좌 연결','로그인 안내','가능','회원 세션과 분리'),('발행된 공통 브리핑 조회','가능','가능','운영은 별도 세션'),('AI 창 열기 / 질문 전송','안내 / 불가','가능 / 가능','회원 세션과 분리'),('평가 등록','불가','가능','운영 집계 조회'),('관리자 라우트','관리자 로그인','관리자 로그인','가능')],[5.1,3.05,3.05,3.45],'접근 가드·spec-v1.2-gap.md FS045·FS066-01. 실제 권한 smoke는 captures/auth-proof.json')


def intersects(a,b,pad=0):
    return a[0]<b[2]+pad and a[2]>b[0]-pad and a[1]<b[3]+pad and a[3]>b[1]-pad


def shot(slide,record,image,x,y,w,h):
    slide.shapes.add_picture(str(CAP/image),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    scale=w/record['viewport']['width']
    d=32*scale
    targets=[(x+m['box']['x']*scale,y+m['box']['y']*scale,x+(m['box']['x']+m['box']['width'])*scale,y+(m['box']['y']+m['box']['height'])*scale) for m in record.get('markers',[])]
    circles=[]
    for m,b in zip(record.get('markers',[]),targets):
        outline=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(b[0]),Inches(b[1]),Inches(b[2]-b[0]),Inches(b[3]-b[1]))
        outline.fill.background(); outline.line.color.rgb=BRAND; outline.line.width=Pt(2*scale*72)
        for effect in outline._element.xpath('./p:style/a:effectRef'): effect.set('idx','0')
        gap=max(.03,3*scale)
        candidates=[(b[0]-d/2-gap,b[1]-d/2-gap),(b[2]+d/2+gap,b[1]-d/2-gap),(b[0]-d/2-gap,(b[1]+b[3])/2),(b[2]+d/2+gap,(b[1]+b[3])/2),(b[0]-d/2-gap,b[3]+d/2+gap)]
        for distance in range(1,12):
            candidates.extend([(b[0]-d/2-gap,b[1]-d/2-gap-distance*d),(b[2]+d/2+gap,b[1]-d/2-gap-distance*d)])
        chosen=None
        for cx,cy in candidates:
            area=(cx-d/2,cy-d/2,cx+d/2,cy+d/2)
            if area[0]<.12 or area[1]<1.32 or area[2]>11.3 or area[3]>8.32: continue
            if any(intersects(area,other,.02) for other in circles+targets): continue
            chosen=(cx,cy,area);break
        if chosen is None: raise RuntimeError(f'번호 원을 비킬 자리가 없음: {record["id"]} {m["number"]}')
        cx,cy,area=chosen;circles.append(area)
        if abs(cx-(b[0]-d/2))>.01 or abs(cy-(b[1]-d/2))>.01:
            leader=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(cx),Inches(cy),Inches(b[0]),Inches(b[1]))
            leader.line.color.rgb=BRAND;leader.line.width=Pt(.7)
        circle=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(area[0]),Inches(area[1]),Inches(d),Inches(d))
        circle.fill.solid();circle.fill.fore_color.rgb=BRAND;circle.line.color.rgb=C('text-inverse');circle.line.width=Pt(2*scale*72)
        for effect in circle._element.xpath('./p:style/a:effectRef'): effect.set('idx','0')
        tf=circle.text_frame;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        p=tf.paragraphs[0];p.text=str(m['number']);p.alignment=PP_ALIGN.CENTER;p.font.name=FONT;p.font.bold=True;p.font.size=Pt(d*72*.5);p.font.color.rgb=C('text-inverse')
        annotation_geometry.append({'id':record['id'],'number':m['number'],'image':image,'target':list(b),'circle':list(area)})


annotation_geometry=[]
from PIL import ImageFont
measure_font=ImageFont.truetype(str(ROOT/'prototype/web/node_modules/@sun-typeface/suit/fonts/static/ttf/SUIT-Regular.ttf'),1000)
def text_lines(text,size,width):
    result=0
    for paragraph in text.split('\n'):
        used=0; count=1
        for char in paragraph:
            advance=measure_font.getlength(char)*size/1000
            if used+advance>(width-.06)*72:
                count+=1;used=0
            used+=advance
        result+=count
    return result

for item in manifest:
    s=new(item['id'],item['title'],item['route'])
    all_markers=list(item['markers'])
    if item.get('pair'):
        box(s,.8,1.25,3.15,.25,'누르기 전',11,True); box(s,5.4,1.25,3.15,.25,'누른 뒤',11,True)
        shot(s,item,item['image'],.8,1.52,3.05,6.60)
        after=json.loads((CAP/item['pair'].replace('.png','.json')).read_text())
        shot(s,after,item['pair'],5.4,1.52,3.05,6.60)
        all_markers.extend(after['markers'])
    else:
        shot(s,item,item['image'],.4,1.45,10.8,6.75)
    state_refs='코드 근거: '+item['stateSource']+'\n'+'\n'.join('기존 캡처: '+p for p in item['stateImages'])
    sections=[('진입 조건',item['entry']),('누르면','\n'.join(f"{m['number']}  {m['result']} → {m['destination']}" for m in all_markers)),('상태','현재 캡처: '+item['capturedState']+'\n구현 상태: '+item['states']+'\n'+state_refs),('명세',' · '.join(item['fs'].split('·'))+' / 요구 '+(item['requirements'] or '미기재'))]
    if item['id'].startswith('D-SEN-'): sections.insert(0,('예시·데이터 한계',item['note']))
    elif item['id'] in ('D-GRP-02','D-GRP-03','D-COM-03','D-COM-04','D-COM-05','D-COM-06','A-ADM-04b'): sections.append(('조회·선택 규칙',item['note']))
    # 반복 범위 고지는 발표자 노트에 둔다. 모의·근거·시점 고지는 실제 화면과 공통 규칙에 있다.
    team={
        'D-AUT-02':'가입 필드와 약관 본문을 확정해야 해요.',
        'D-AUT-03':'실제 제공자와 계정 연결 정책을 확정해야 해요.',
        'D-GRP-01b':'필터 깊이 1단계는 프로토타입 가정이에요.',
        'D-GRP-06':'영향 경로라는 이름과 중간 산업·제품 자료를 확인해야 해요.',
        'D-BRF-01':'지연 판정 수치·환율 비교 기준·보존 기간이 미정이에요.',
        'D-BRF-02':'지연 판정 수치·환율 비교 기준·보존 기간이 미정이에요.',
        'A-ADM-02b':'대화록의 별도 권한·열람 목적·실패 차단이 미정이에요.',
        'A-ADM-03':'R1 승인·검수 권한·허용 후 재검증 방식이 미정이에요.',
        'A-ADM-05':'관리자 권한 등급과 회수 정책이 미정이에요.',
    }.get(item['id'])
    if team: sections.append(('팀 확인 필요',team))
    side_x=11.55; side_w=4.05
    for body_size in (12.5,12,11.5):
        heights=[text_lines(body,body_size if title!='상태' else body_size-1,side_w)*((body_size if title!='상태' else body_size-1)*1.30/72)+.10 for title,body in sections]
        total=sum(heights)+len(sections)*.32
        if total<=6.75: break
    if total>6.75: raise RuntimeError(f'설명 열이 너무 김: {item["id"]} {total}')
    y=1.45
    for (title,body),height in zip(sections,heights):
        box(s,side_x,y,side_w,.23,title,13.5,True,color=BRAND,spacing=0,line_spacing=1.0);y+=.27
        size=body_size if title!='상태' else body_size-1
        box(s,side_x,y,side_w,height,body,size,spacing=0,line_spacing=Pt(size*1.30));y+=height+.05
    s.notes_slide.notes_text_frame.text='캡처 출처: '+item['image']+'\n좌표·클릭 후 실제 화면: '+item['image'].replace('.png','.json')+'\n'+item['visibleText']

# 부록은 정책 문서와 실제 캡처 범위를 명확히 구분한다.
table('APP-EVIDENCE','근거·원문 접근 상태 매트릭스',['상태','화면 표시','다음 행동'],[
('원문 링크 있음','인용·출처 + 원문 링크','새 창. 기존 패널·대화 유지'),('원문 링크 없음','원문을 제공하지 않는 자료입니다','인용·출처 유지'),('원문 접근 실패','앱이 외부 사이트 접근을 검사하지 않음','외부 사이트에서 처리. 앱 오류 만들지 않음'),('연결 자료 0건','근거 없음. 근거 버튼 생략','빈 상태, 오류 아님'),('근거 조회 일시 실패','받은 근거 유지 + 실패 건수','실패한 근거만 다시 시도'),('삭제·비공개·권한 없음','제공 불가. 보호 내용 표시 안 함','404·403은 재시도 없음'),('이전 대상 늦은 응답','현재 대상에 섞이지 않음','대상별 조회 칸 분리')],[3.25,6.15,5.25],'design-direction.md §5.5. 전 상태 실측 아님 · FS024·FS050-01·FS043')
table('APP-ID','팀 v1.7 ↔ v2.1 화면 ID 대응',['팀 v1.7','프로토타입 v2.1','판정·차이'],[
('D-GRP-05 필터','D-GRP-01 · D-GRP-01b','두 모드와 개인화 내부 필터로 분리'),
('D-GRP-02 동시 이슈','D-GRP-02 사건 패널','명칭·단위 변경'),
('D-GRP-03 영향 경로','D-GRP-03 기업 패널','기업 정보 패널과 경로 역할 분리'),
('D-GRP-05 관계 설명 · D-GRP-06 간선 모달','D-GRP-04 연결선 상세','사건 공유·후속 사건 상세로 한정'),
('D-GRP-04 근거','D-GRP-05 근거 창','ID 이동'),
('D-GRP-03 영향 경로 · D-GRP-06 간선 모달','D-GRP-06 영향 경로','기업 패널의 관련 사건에서 진입'),
('D-BRF-03 장 마감 보드','D-BRF-02 퇴근 회차 비교','독립 화면 없이 흡수'),
('지도 안 목록 보기','D-STK-01 · M-STK-01','독립 종목 메뉴 /companies로 분리'),
('A-ADM-01~05','관리자 로그인 → 최종 5개 묶음','42쪽 기준, 28쪽의 이전 분류는 승계하지 않음')],[4.6,5.3,4.75],'팀 원문 screen-plan-v1.7.txt · 본문과 같은 화면 흐름 순서')
table('APP-VERDICT-01','팀 v1.7 판정표 · 채택',['팀 항목','v2.0 적용','근거'],[
('D-BRF-03 회차 비교','D-BRF-02에 흡수','공통 발행본을 날짜·출근/퇴근으로 비교'),
('모바일 아침 이후 바뀐 내용','M-BRF-01 · D-BRF-02','동일 날짜 출근/퇴근 이슈 ID 비교'),
('근거·원문 상태 매트릭스','APP-EVIDENCE','일시 실패와 404·403 제공 불가 분리'),
('인과관계 아님 고지','D-COM-01b','명세 FS040 · 10-03 시점 겹침 고지'),
('모바일 하단 AI 탭','M-AI-01','새 화면 없이 기존 질문 창 열기'),
('권한표·관리자 감사 형식','AUTHORITY · A-ADM-05b','권한과 변경 기록을 화면에서 구분')],[4.7,4.25,5.7],'작성 명세 팀 v1.7 확정 판정 · review.md Main 정정')
table('APP-VERDICT-02','팀 v1.7 판정표 · 불채택',['팀 항목','v2.0 판정','근거'],[
('이슈–기업/기업 간 간선 모달 D-GRP-06','모달 대신 영향 경로·연결선 패널','FS022·D03. 사건–기업 선 클릭 제거'),
('외부 원인 10종·공급/경쟁 직결선','기업·사건 두 노드와 실제 관계만','값을 만드는 데이터 없음'),
('긍정 가능/부정 가능 배지','방향 값 생기면 4분류 규격 적용','명세는 방향을 요구하나 PoC 생산자 없음. 문구 미정'),
('모듈별 6색·카드 좌측 색 막대','중성 회색 + 의미별 색','브랜드 주 행동·선택, 모듈 ID·아이콘'),
('관심 최대 3개','무제한·0개 허용','9/30 결정·FS012·FS027'),
('브리핑 08:00/18:00','07:40/18:10 생성 시작','10-01 명세 정본')],[4.7,4.25,5.7],'작성 명세 팀 v1.7 확정 판정 · review.md Main 정정')
columns('APP-POLICY-01','팀 확인 필요 · 관리자와 운영','팀에 요청한 결정 1~3','1. R1을 자동 검증 정책의 일부로 인정할지\n역할이 확인되지 않은 후보를 역할 없이 공개해요.\n2. 관리자 권한 등급·대화록 범위·마스킹\n별도 열람 권한·목적·실패 차단은 남은 충돌이에요.\n3. 관리자 허용 → 재검증 전환\n현재 즉시 공개와 FS067-02가 충돌해요.','팀에 요청한 결정 4~6','4. 운영 값: 로그인 실패 제한·재시도 범위·보존 기간\nAI 재시도 3회·중단 120초는 프로토타입 기본값\n5. 그래프 기본값: 사이드바·패널·원 최소 크기\n영향 경로라는 이름의 오해 가능성 확인\n6. 시세 공급원: 토스 API 이용 목적·제3자 제공\n키 발급과 별개로 약관·배포 권한 확인','progress-2026-10-02.md §1. 여섯 항목은 확정 결정이 아니라 결정 요청이에요.')
columns('APP-POLICY-02','팀 확인 필요 · 10-03 명세 충돌','명세 갱신본','FS017: 기존 시트 제품 노드 vs v2 기업·사건 두 노드\n현재 캡처는 기업·사건만. 정본 대체 선언은 없어요.\nFS004 vs FS061: 제공자·신규 계정·연결 정책 정합화\n현재는 /mock-idp 모의 제공자예요.\nFS014: 초기 그래프에 튜토리얼 기준 혼입\n초기 설정·튜토리얼 계약을 분리해야 해요.','Notion 최종 시퀀스 vs 9/30','비회원 관심 선택·저장 금지 vs 임시 저장 허용\n관심 무제한·0개 허용 vs 최소1/최대3\n사건–기업 간선 클릭 제거 vs 상세 패널 진입\n보유·관심 한 모드 내 필터 vs 세 모드 혼재\n가입 자동 로그인 vs 로그인 화면으로 복귀\n문서 제목 최종만으로 팀 승인으로 보지 않아요.','spec-wbs-delta-2026-10-03.md §6 · notion-delta-2026-10-03.md §7')
columns('APP-LIMIT','근거 버전·남은 한계','현재 근거','기능명세: functional-spec_2026-10-06\n요구사항: requirements-sheet_2026-10-06\nWBS: wbs_2026-10-03(내용 동일)\n변경 대응: doc-delta-2026-10-06.md §2.2\n화면: captures/ 원본·좌표·클릭 결과\n디자인: design-direction.md · sentiment.md','정책·데이터 한계','시장 요약·재무 추이·산업/테마 생산자 없음\n데모 판정값 없음. 예시 화면과 실제 데이터 분리\n생성형 AI·실 소셜 IdP·정기 수집원 미연결\n환율 변화 기준·지연 수치는 프로토타입 가정\n발행 뒤 지표 재조회는 회차 불변 정책과 충돌\n실제 오류·권한·외부 원문 실패 전수 캡처는 아님','없는 데이터·전망·확률을 만들어 채우지 않아요.')
table('APP-DOCS','문서 체계 · 화면에서 구현 추적까지',['문서','확인할 내용','저장소 위치'],[
('화면계획서','실제 화면·번호 조작·상태·화면 ID','docs/screen-plan/zoop-screen-plan-v2.1.pdf'),
('기능 명세','10-06 고정본의 FS·NFS 원문','docs/evidence/raw/functional-spec_2026-10-06/'),
('목표 API 명세','기능·화면별 API 계약과 BFF 차이','docs/api/out/zoop-api-spec.md'),
('WBS','업무 번호·명세 연결·담당·일정','docs/evidence/raw/wbs_2026-10-03/'),
('통합 추적표','화면 ↔ 기능 ↔ API ↔ WBS ↔ 구현 판정','docs/system/out/zoop-trace.xlsx'),
('통합 문서 목록','화면별·기능별 문서와 검사 결과','docs/system/out/index.md')],[2.5,5.2,6.95],'docs/system/README.md · links.yaml · API 정본 openapi.yaml')

pptx=BASE/'zoop-screen-plan-v2.1.pptx'; prs.save(pptx)
# SUIT가 시스템에 없어도 편집·PDF 변환에 쓰도록 OFL 글꼴을 PPTX 안에 넣는다.
# EOT 0x00010000 구조: https://www.w3.org/submissions/EOT/Overview.html §3.1
import struct, zipfile
from lxml import etree
def eot(path):
    data=path.read_bytes()
    count=struct.unpack_from('>H',data,4)[0]
    tables={}
    for n in range(count):
        tag,checksum,offset,length=struct.unpack_from('>4sIII',data,12+n*16)
        tables[tag]=data[offset:offset+length]
    os2=tables[b'OS/2']; head=tables[b'head']
    weight=struct.unpack_from('>H',os2,4)[0]; fst=struct.unpack_from('>H',os2,8)[0]
    assert not fst & 2, '글꼴이 임베딩을 허용하지 않음'
    ranges=struct.unpack_from('>4I',os2,42); pages=struct.unpack_from('>2I',os2,78)
    header=struct.pack('<IIII10sBBIHH4I2II4I',0,len(data),0x10000,0,os2[32:42],1,0,weight,fst,0x504c,*ranges,*pages,struct.unpack_from('>I',head,8)[0],0,0,0,0)
    style='Bold' if weight>=700 else 'Regular'
    for name in ('SUIT',style,'Version 2.0','SUIT '+style):
        b=name.encode('utf-16-le'); header+=struct.pack('<HH',0,len(b))+b
    payload=header+data
    return struct.pack('<I',len(payload))+payload[4:]
with zipfile.ZipFile(pptx) as z:
    parts={name:z.read(name) for name in z.namelist()}
P='http://schemas.openxmlformats.org/presentationml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
presentation=etree.fromstring(parts['ppt/presentation.xml']); presentation.set('embedTrueTypeFonts','1'); presentation.set('saveSubsetFonts','0')
fontlist=etree.SubElement(presentation,f'{{{P}}}embeddedFontLst'); font=etree.SubElement(fontlist,f'{{{P}}}embeddedFont')
defaultstyle=presentation.find(f'{{{P}}}defaultTextStyle')
if defaultstyle is not None:
    presentation.remove(fontlist); presentation.insert(presentation.index(defaultstyle),fontlist)
etree.SubElement(font,f'{{{P}}}font',typeface='SUIT',pitchFamily='34',charset='1')
rels=etree.fromstring(parts['ppt/_rels/presentation.xml.rels'])
ct=etree.fromstring(parts['[Content_Types].xml'])
etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='fntdata',ContentType='application/x-fontdata')
etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='txt',ContentType='text/plain')
parts['ppt/fonts/LICENSE.txt']=(ROOT/'prototype/web/node_modules/@sun-typeface/suit/LICENSE').read_bytes()
fontdir=ROOT/'prototype/web/node_modules/@sun-typeface/suit/fonts/static/ttf'
for n,style in enumerate(('Regular','Bold'),1):
    rid=f'rIdSuit{n}'
    etree.SubElement(font,f'{{{P}}}{style.lower()}',{f'{{{R}}}id':rid})
    etree.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id=rid,Type=R+'/font',Target=f'fonts/suit{n}.fntdata')
    parts[f'ppt/fonts/suit{n}.fntdata']=eot(fontdir/f'SUIT-{style}.ttf')
parts['ppt/presentation.xml']=etree.tostring(presentation)
parts['ppt/_rels/presentation.xml.rels']=etree.tostring(rels)
(CAP/'annotation-geometry.json').write_text(json.dumps(annotation_geometry,ensure_ascii=False,indent=2))
parts['[Content_Types].xml']=etree.tostring(ct)
with zipfile.ZipFile(pptx,'w',zipfile.ZIP_DEFLATED) as z:
    for name,data in parts.items(): z.writestr(name,data)
(CAP/'slide-index.json').write_text(json.dumps(index,ensure_ascii=False,indent=2))
print('PPTX',len(prs.slides),pptx)
parser=argparse.ArgumentParser(); parser.add_argument('--pdf',action='store_true'); args=parser.parse_args()
if args.pdf:
    fonts=ROOT/'prototype/web/node_modules/@sun-typeface/suit/fonts/static/ttf'
    subprocess.run(['/Applications/LibreOffice.app/Contents/MacOS/soffice','-env:UserInstallation=file:///tmp/zoop-screen-plan-lo','--headless','--convert-to','pdf:impress_pdf_Export:{"EmbedStandardFonts":{"type":"boolean","value":"true"}}','--outdir',str(BASE),str(pptx)],env=os.environ|{'SAL_FONTPATH':str(fonts)},check=True)
    import pymupdf as fitz
    from PIL import Image, ImageDraw
    pdf=fitz.open(BASE/'zoop-screen-plan-v2.1.pdf')
    render=CAP/'rendered'
    if render.exists(): shutil.rmtree(render)
    render.mkdir()
    audit=[]
    for n,page in enumerate(pdf,1):
        pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)); pix.save(render/f'page-{n:02}.png')
        fonts_used=page.get_fonts(full=True)
        audit.append({'page':n,'id':index[n-1]['id'],'fonts':[{'font':f[3],'embedded':bool(pdf.extract_font(f[0])[3])} for f in fonts_used]})
    (render/'font-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
    # 한 장씩 확인할 수 있는 원본과 6장씩 모은 검토용 이미지를 함께 보존한다.
    for start in range(0,len(pdf),6):
        contact=Image.new('RGB',(1800,1100),'white'); draw=ImageDraw.Draw(contact)
        for k,n in enumerate(range(start,min(start+6,len(pdf)))):
            im=Image.open(render/f'page-{n+1:02}.png'); im.thumbnail((880,495))
            x=(k%2)*900; y=(k//2)*366
            # 16:9 전체를 630×354로 배치해 검토 번호와 함께 출력한다.
            im.thumbnail((630,354)); contact.paste(im,(x,y)); draw.text((x+640,y+20),f'{n+1:02} {index[n]["id"]}',fill='black')
        contact.save(render/f'contact-{start//6+1:02}.png')
    outside=[]; overlaps=[]; fallbacks=[]
    for number,page in enumerate(pdf,1):
        lines=[]
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines',[]):
                text=''.join(span['text'] for span in line['spans']).strip()
                if not text: continue
                rect=fitz.Rect(line['bbox'])
                if not page.rect.contains(rect): outside.append({'page':number,'text':text,'box':list(rect)})
                lines.append((text,rect))
        for n,(text,rect) in enumerate(lines):
            for other,area in lines[n+1:]:
                cross=rect & area
                if cross.width<=1 or cross.height<=1: continue
                issue={'page':number,'first':text,'second':other,'intersection':list(cross)}
                if text in ('§','–','·') or other in ('§','–','·'): fallbacks.append(issue)
                else: overlaps.append(issue)
    density={}; references=0
    for item in manifest:
        records=[item]
        if item.get('pair'): records.append(json.loads((CAP/item['pair'].replace('.png','.json')).read_text()))
        density[item['id']]=sum(len(record['markers']) for record in records)
        for record in records:
            assert len(record['markers'])==len(record['clickEvidence']), record['id']
            for evidence in record['clickEvidence']:
                if 'reference' in evidence:
                    assert 'verifiedInteraction' in evidence, (record['id'],evidence)
                    references+=1
    assert all(3<=n<=6 for n in density.values()), density
    checks={'pdfPages':len(pdf),'pptxSlides':len(prs.slides),'captureGroups':len(manifest),
        'numberedControls':sum(density.values()),'densityById':density,'confirmedNavigationReferences':references,
        'annotationBoxesAndCirclesMatch':len(annotation_geometry),'textOutsidePage':outside,
        'possibleTextOverlaps':overlaps,'fontFallbackSpanIntersections':fallbacks,
        'allFontsEmbedded':all(font['embedded'] for page in audit for font in page['fonts']),
        'visualReview':'미기록: page-*.png 전 페이지를 직접 검토한 뒤 결과를 기록한다.'}
    (CAP/'verification.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n')
    print('AUDIT',sum(density.values()),'controls;',len(outside),'outside;',len(overlaps),'possible overlaps')
    print('PDF',len(pdf),'pages; rendered',render)
