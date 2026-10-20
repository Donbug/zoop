"""Capture the public functional-spec XLSX export for 2026-10-18; immutable baselines, no credentials or installs.

Copy of refresh_2026-10-12/_capture_workbooks.py with only the date, sources and baselines changed: fetches the
functional-spec workbook (wbs_2026-10-12 and requirements-sheet_2026-10-06 stay the carried freezes), compares it
with the 2026-10-12 capture and attributes every changed cell to _logged-changes.json or to an unlogged editor.
"""
import csv
import hashlib
import io
import json
import re
import sys
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

DATE = "2026-10-18"
ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
REPO = ROOT.parent.parent.parent
SOURCES = [
    ("functional-spec", "functional-spec_2026-10-12", "1irLUY-NCkDRfF8fO3WkZWU2MpaStu2RPmQ-wG0F9Xcs"),
]
CARRIED = {"wbs": "docs/evidence/raw/wbs_2026-10-12", "requirements-sheet": "docs/evidence/raw/requirements-sheet_2026-10-06"}
NEW_DIRS = {OUT} | {ROOT / f"{key}_{DATE}" for key, _, _ in SOURCES}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def rel(path):
    return str(path.relative_to(REPO))


def write_new(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as target:
        target.write(data)


def save_json(path, value):
    write_new(path, (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))


def strings(worksheet):
    return [["" if value is None else str(value) for value in row]
            for row in worksheet.iter_rows(values_only=True)]


def hyperlinks(workbook):
    return {(ws.title, cell.coordinate): {
        "target": cell.hyperlink.target, "location": cell.hyperlink.location,
        "display": cell.hyperlink.display, "tooltip": cell.hyperlink.tooltip}
        for ws in workbook.worksheets for row in ws for cell in row if cell.hyperlink}


def fetch(source):
    key, _, source_id = source
    url = f"https://docs.google.com/spreadsheets/d/{source_id}/export?format=xlsx"
    started = datetime.now(timezone.utc)
    request = urllib.request.Request(url, headers={"User-Agent": "Zoop-Public-Evidence-Recheck/1.0"})
    with urllib.request.urlopen(request, timeout=90) as response:
        data = response.read()
        status = response.status
        content_type = response.headers.get("Content-Type")
    completed = datetime.now(timezone.utc)
    if status != 200 or not data.startswith(b"PK"):
        raise RuntimeError(f"{key}: expected public XLSX HTTP 200, got {status}, {content_type}")
    return source, data, {
        "source_url": f"https://docs.google.com/spreadsheets/d/{source_id}", "export_url": url,
        "request_started_at_utc": started.isoformat(), "collected_at_utc": completed.isoformat(),
        "collected_at_kst": completed.astimezone(ZoneInfo("Asia/Seoul")).isoformat(),
        "auth": "none", "http_status": status, "content_type": content_type,
        "response_bytes": len(data), "response_sha256": sha(data),
    }


def verify_csv(xlsx_path, sheet_files):
    workbook = load_workbook(xlsx_path, data_only=True)
    if list(sheet_files) != workbook.sheetnames:
        raise AssertionError((str(xlsx_path), "sheet names/order"))
    records = []
    for worksheet in workbook.worksheets:
        csv_path = xlsx_path.parent / sheet_files[worksheet.title]
        with csv_path.open(encoding="utf-8-sig", newline="") as source:
            actual = list(csv.reader(source))
        if len(actual) != worksheet.max_row:
            raise AssertionError((str(csv_path), "row count"))
        for row_number in range(1, worksheet.max_row + 1):
            if len(actual[row_number - 1]) != worksheet.max_column:
                raise AssertionError((str(csv_path), row_number, "column count"))
            for column_number in range(1, worksheet.max_column + 1):
                value = worksheet.cell(row_number, column_number).value
                expected = "" if value is None else str(value)
                if actual[row_number - 1][column_number - 1] != expected:
                    raise AssertionError((str(csv_path), row_number, column_number))
        records.append({"sheet": worksheet.title, "csv": rel(csv_path),
                        "rows": worksheet.max_row, "columns": worksheet.max_column,
                        "cells": worksheet.max_row * worksheet.max_column,
                        "sha256": sha(csv_path.read_bytes()), "result": "passed"})
    workbook.close()
    return records


def attribute(deltas, logged, hyperlink_deltas, structures):
    """Split deltas into logged-and-matching, logged-but-mismatched and unlogged (other editor) changes.

    A logged ISO date such as 2026-10-12 is accepted as the Sheets date cell it becomes
    (openpyxl reads it as 2026-10-12 00:00:00); a logged URL cell is accepted as the hyperlink
    Sheets auto-attaches to it. Both are recorded explicitly, never silently.
    """
    expected = {}
    for change in logged["changes"]:
        for cell, values in change["cells"].items():
            expected[(change["sheet"], cell)] = (change["id"], values)
    seen, matched, mismatched, unlogged = set(), [], [], []
    for delta in deltas:
        key = (delta["sheet"], delta["cell"])
        if key not in expected:
            unlogged.append(delta)
            continue
        change_id, values = expected[key]
        seen.add(key)
        record = {"change": change_id, "workbook": delta["workbook"], "sheet": key[0], "cell": key[1],
                  "before": delta["before"], "after": delta["after"]}
        typed_date = re.fullmatch(r"\d{4}-\d{2}-\d{2}", values["after"]) is not None and delta["after"] == values["after"] + " 00:00:00"
        if delta["before"] == values["before"] and (delta["after"] == values["after"] or typed_date):
            matched.append(dict(record, typing="sheets-date-cell" if typed_date else "exact"))
        else:
            mismatched.append(dict(record, expected_before=values["before"], expected_after=values["after"]))
    absent = [{"change": expected[key][0], "sheet": key[0], "cell": key[1]}
              for key in expected if key not in seen]
    links_logged, links_unlogged = [], []
    for delta in hyperlink_deltas:
        key = (delta["sheet"], delta["cell"])
        target = (delta["after"] or {}).get("target")
        if delta["before"] is None and key in expected and target == expected[key][1]["after"]:
            links_logged.append(dict(delta, change=expected[key][0], typing="sheets-autolink-of-logged-url"))
        else:
            links_unlogged.append(delta)
    explained = []
    for structure in structures:
        before, after = structure["before"], structure["after"]
        rows = range(before[0] + 1, after[0] + 1) if before and after and before[1] == after[1] else range(0)
        cell_rows = {int(re.sub(r"\D", "", cell)) for sheet, cell in expected if sheet == structure["sheet"]}
        explained.append(dict(structure, explained_by_logged_cells=bool(rows) and all(row in cell_rows for row in rows)))
    return matched, mismatched, absent, unlogged, links_logged, links_unlogged, explained


baseline_hashes = {path: sha(path.read_bytes()) for path in sorted(ROOT.rglob("*"))
                   if path.is_file() and not any(parent in NEW_DIRS for parent in path.parents)}
logged = json.loads((OUT / "_logged-changes.json").read_text(encoding="utf-8"))
for key, _, _ in SOURCES:
    if (ROOT / f"{key}_{DATE}").exists():
        raise RuntimeError(f"{key}_{DATE} already exists; refusing to overwrite a frozen capture")
with ThreadPoolExecutor(max_workers=len(SOURCES)) as pool:
    fetched = list(pool.map(fetch, SOURCES))

all_deltas, all_formula_deltas, all_mapping_deltas = [], [], []
results, verification = [], []
for (key, baseline_name, _), data, metadata in fetched:
    baseline_dir = ROOT / baseline_name
    baseline_manifest = json.loads((baseline_dir / "_collection-manifest.json").read_text())
    baseline_path = baseline_dir / baseline_manifest["xlsx"]
    old = load_workbook(baseline_path, data_only=True)
    new = load_workbook(io.BytesIO(data), data_only=True)
    old_raw = load_workbook(baseline_path, data_only=False)
    new_raw = load_workbook(io.BytesIO(data), data_only=False)
    before_links, after_links = hyperlinks(old_raw), hyperlinks(new_raw)
    hyperlink_deltas = [{"workbook": key, "sheet": sheet, "cell": cell,
                        "before": before_links.get((sheet, cell)), "after": after_links.get((sheet, cell))}
                       for sheet, cell in sorted(set(before_links) | set(after_links))
                       if before_links.get((sheet, cell)) != after_links.get((sheet, cell))]
    states_before = {ws.title: ws.sheet_state for ws in old.worksheets}
    states_after = {ws.title: ws.sheet_state for ws in new.worksheets}
    deltas, formula_deltas, structures, mapping_deltas = [], [], [], []
    for sheet in dict.fromkeys(old.sheetnames + new.sheetnames):
        bw = old[sheet] if sheet in old.sheetnames else None
        aw = new[sheet] if sheet in new.sheetnames else None
        before = strings(bw) if bw else []
        after = strings(aw) if aw else []
        before_shape = [bw.max_row, bw.max_column] if bw else None
        after_shape = [aw.max_row, aw.max_column] if aw else None
        if before_shape != after_shape:
            structures.append({"sheet": sheet, "before": before_shape, "after": after_shape})
        changed_rows = set()
        for row in range(1, max(len(before), len(after)) + 1):
            b = before[row - 1] if row <= len(before) else []
            a = after[row - 1] if row <= len(after) else []
            for column in range(1, max(len(b), len(a)) + 1):
                bv = b[column - 1] if column <= len(b) else ""
                av = a[column - 1] if column <= len(a) else ""
                if bv != av:
                    deltas.append({"workbook": key, "sheet": sheet, "row": row, "column": column,
                                   "cell": f"{get_column_letter(column)}{row}", "before": bv, "after": av})
                    changed_rows.add(row)
        br = old_raw[sheet] if sheet in old_raw.sheetnames else None
        ar = new_raw[sheet] if sheet in new_raw.sheetnames else None
        formulas_before = {cell.coordinate: str(cell.value) for row in br for cell in row if cell.data_type == "f"} if br else {}
        formulas_after = {cell.coordinate: str(cell.value) for row in ar for cell in row if cell.data_type == "f"} if ar else {}
        for cell in sorted(set(formulas_before) | set(formulas_after)):
            bv, av = formulas_before.get(cell, ""), formulas_after.get(cell, "")
            if bv != av:
                formula_deltas.append({"workbook": key, "sheet": sheet, "cell": cell, "before": bv, "after": av})
        for row in sorted(changed_rows):
            b = before[row - 1] if row <= len(before) else []
            a = after[row - 1] if row <= len(after) else []
            row_changes = [item for item in deltas if item["sheet"] == sheet and item["row"] == row]
            fs_matches = sorted(set(re.findall(r"\b(?:N?FS)\d+(?:-\d+)?\b", "\n".join(b + a))))
            if fs_matches or sheet == "명세 목차" or (key == "wbs" and sheet == "프로젝트 WBS"):
                mapping_deltas.append({"workbook": key, "sheet": sheet, "row": row,
                    "fs_nfs_ids_in_row": fs_matches,
                    "before_nonempty_cells": {get_column_letter(i + 1): v for i, v in enumerate(b) if v},
                    "after_nonempty_cells": {get_column_letter(i + 1): v for i, v in enumerate(a) if v},
                    "cell_deltas": row_changes})
    state_deltas = [{"sheet": sheet, "before": states_before.get(sheet), "after": states_after.get(sheet)}
                    for sheet in dict.fromkeys(list(states_before) + list(states_after))
                    if states_before.get(sheet) != states_after.get(sheet)]
    changed = bool(deltas or formula_deltas or hyperlink_deltas or structures or state_deltas
                   or old.sheetnames != new.sheetnames)
    if not changed:
        raise RuntimeError(f"{key}: no change against {baseline_name}; nothing to freeze")
    capture_dir = ROOT / f"{key}_{DATE}"
    xlsx_name = f"zoop-{key}_{DATE}.xlsx"
    capture_path = capture_dir / xlsx_name
    write_new(capture_path, data)
    sheet_files = {}
    for index, worksheet in enumerate(new.worksheets):
        filename = baseline_manifest["sheets"].get(worksheet.title)
        if filename is None:
            filename = f"{index:02d}-" + re.sub(r"[^\w가-힣.-]+", "-", worksheet.title).strip("-") + ".csv"
        if filename in sheet_files.values():
            raise RuntimeError(f"CSV filename collision: {worksheet.title}: {filename}")
        sheet_files[worksheet.title] = filename
        stream = io.StringIO(newline="")
        csv.writer(stream).writerows(strings(worksheet))
        write_new(capture_dir / filename, stream.getvalue().encode("utf-8-sig"))
    matched, mismatched, absent, unlogged, links_logged, links_unlogged, explained = attribute(
        deltas, {"changes": [c for c in logged["changes"] if c["workbook"] == key]}, hyperlink_deltas, structures)
    manifest = dict(metadata, baseline=baseline_name, baseline_xlsx_sha256=sha(baseline_path.read_bytes()),
        content_identical=False, meaningful_cell_content_changed=bool(deltas),
        data_only=True, csv_encoding="utf-8-sig", xlsx=xlsx_name, sha256=sha(data), sheets=sheet_files,
        sheet_ranges={ws.title: {"rows": ws.max_row, "columns": ws.max_column} for ws in new.worksheets},
        sheet_states=states_after, sheet_state_deltas=state_deltas,
        files={p.name: sha(p.read_bytes()) for p in capture_dir.iterdir() if p.is_file()},
        baseline_sheet_order=old.sheetnames, current_sheet_order=new.sheetnames,
        cell_delta_count=len(deltas), formula_delta_count=len(formula_deltas), structural_delta_count=len(structures),
        logged_cell_deltas_matched=len(matched), logged_cell_deltas_mismatched=len(mismatched),
        logged_cells_absent_from_delta=len(absent), unlogged_cell_deltas=len(unlogged),
        logged_hyperlink_deltas=len(links_logged), unlogged_hyperlink_deltas=len(links_unlogged),
        unexplained_structural_deltas=sum(not item["explained_by_logged_cells"] for item in explained))
    save_json(capture_dir / "_collection-manifest.json", manifest)
    save_json(capture_dir / "_cell-delta.json", deltas)
    save_json(capture_dir / "_formula-delta.json", formula_deltas)
    save_json(capture_dir / "_structure-delta.json", structures)
    save_json(capture_dir / "_mapping-relevant-delta.json", mapping_deltas)
    save_json(capture_dir / "_hyperlink-targets.json", {
        "baseline": [{"sheet": s, "cell": c, **v} for (s, c), v in before_links.items()],
        "current": [{"sheet": s, "cell": c, **v} for (s, c), v in after_links.items()], "deltas": hyperlink_deltas})
    all_deltas.extend(deltas)
    all_formula_deltas.extend(formula_deltas)
    all_mapping_deltas.extend(mapping_deltas)
    current_verification = verify_csv(capture_path, sheet_files)
    verification.extend(current_verification)
    results.append(dict(metadata, workbook=key, baseline=rel(baseline_dir), capture=rel(capture_dir),
        recommended_importer_input=rel(capture_dir), xlsx_sha256=sha(data),
        baseline_xlsx_sha256=sha(baseline_path.read_bytes()), sheet_count=len(new.sheetnames),
        sheet_state_deltas=state_deltas, cell_delta_count=len(deltas), formula_delta_count=len(formula_deltas),
        structural_deltas=structures, mapping_relevant_row_count=len(mapping_deltas),
        hyperlink_delta_count=len(hyperlink_deltas), current_hyperlink_count=len(after_links),
        cells_checked=sum(v["cells"] for v in current_verification), current_csv_readback="passed",
        logged_cell_deltas_matched=len(matched), logged_cell_deltas_mismatched=len(mismatched),
        logged_cells_absent_from_delta=len(absent), unlogged_cell_deltas=len(unlogged),
        logged_hyperlink_deltas=len(links_logged), unlogged_hyperlink_deltas=len(links_unlogged),
        unexplained_structural_deltas=sum(not item["explained_by_logged_cells"] for item in explained),
        _attribution={"matched": matched, "mismatched": mismatched, "absent": absent, "unlogged": unlogged,
                      "hyperlinks_logged": links_logged, "hyperlinks_unlogged": links_unlogged,
                      "structures": explained}))
    for workbook in (old, new, old_raw, new_raw):
        workbook.close()

after_hashes = {path: sha(path.read_bytes()) for path in baseline_hashes}
unchanged_originals = after_hashes == baseline_hashes
if not unchanged_originals:
    raise AssertionError("original frozen file bytes changed")
attribution = {"logged_changes_source": rel(OUT / "_logged-changes.json"), "compared_against": "2026-10-12 frozen capture",
    "workbooks": {r["workbook"]: r.pop("_attribution") for r in results}}
attribution["counts"] = {
    "logged_cells_expected": sum(len(c["cells"]) for c in logged["changes"]),
    "logged_cell_deltas_matched": sum(r["logged_cell_deltas_matched"] for r in results),
    "logged_cell_deltas_mismatched": sum(r["logged_cell_deltas_mismatched"] for r in results),
    "logged_cells_absent_from_delta": sum(r["logged_cells_absent_from_delta"] for r in results),
    "unlogged_cell_deltas": sum(r["unlogged_cell_deltas"] for r in results),
    "logged_hyperlink_deltas": sum(r["logged_hyperlink_deltas"] for r in results),
    "unlogged_hyperlink_deltas": sum(r["unlogged_hyperlink_deltas"] for r in results),
    "unexplained_structural_deltas": sum(r["unexplained_structural_deltas"] for r in results),
    "total_cell_deltas": len(all_deltas)}
save_json(OUT / "_change-attribution.json", attribution)
save_json(OUT / "_cell-delta.json", all_deltas)
save_json(OUT / "_formula-delta.json", all_formula_deltas)
save_json(OUT / "_mapping-relevant-delta.json", all_mapping_deltas)
summary = {"result": "passed", "method": {"python": sys.version, "openpyxl": openpyxl.__version__,
    "data_only": True, "csv_encoding": "utf-8-sig", "empty_cell": "", "other_values": "str(value)",
    "all_worksheets": True, "comparison": "exact string cell values, sheet names/order, rectangular ranges and sheet visibility; formula and hyperlink target deltas recorded",
    "network": "one unauthenticated public XLSX GET request (functional-spec); wbs and requirements-sheet not re-fetched; no external writes; no package installs"},
    "workbooks": results, "carried_unchanged": CARRIED,
    "counts": {"workbooks": len(results), "sheets": len(verification),
    "current_cells_checked": sum(v["cells"] for v in verification), "cell_deltas": len(all_deltas),
    "formula_deltas": len(all_formula_deltas), "mapping_relevant_rows": len(all_mapping_deltas),
    **attribution["counts"]},
    "invariants": {"all_tabs_captured": True, "full_workbook_to_csv_readback": "passed",
    "original_frozen_files_unchanged": unchanged_originals, "original_files_hash_checked": len(baseline_hashes)},
    "recommended_importer_inputs": {**{v["workbook"]: v["recommended_importer_input"] for v in results}, **CARRIED}}
save_json(OUT / "_source-recheck.json", summary)
save_json(OUT / "_csv-verification.json", {"result": "passed", "sheets": verification,
    "cells": sum(v["cells"] for v in verification)})
print(json.dumps(summary, ensure_ascii=False, indent=2))
