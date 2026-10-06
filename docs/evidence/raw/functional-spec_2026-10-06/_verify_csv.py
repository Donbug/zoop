"""독립 검증: 고정 xlsx의 저장된 계산값과 CSV의 모든 셀·빈 행을 대조한다."""
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
results = []
for directory_name in ("functional-spec_2026-10-06", "requirements-sheet_2026-10-06"):
    directory = ROOT / directory_name
    manifest = json.loads((directory / "_collection-manifest.json").read_text())
    workbook = load_workbook(directory / manifest["xlsx"], data_only=True)
    for worksheet in workbook.worksheets:
        csv_path = directory / manifest["sheets"][worksheet.title]
        with csv_path.open(encoding="utf-8-sig", newline="") as source:
            actual = list(csv.reader(source))
        assert len(actual) == worksheet.max_row, (csv_path.name, "row count")
        checked = 0
        for row_number in range(1, worksheet.max_row + 1):
            assert len(actual[row_number - 1]) == worksheet.max_column, (csv_path.name, row_number, "column count")
            for column_number in range(1, worksheet.max_column + 1):
                value = worksheet.cell(row_number, column_number).value
                expected = "" if value is None else str(value)
                assert actual[row_number - 1][column_number - 1] == expected, (csv_path.name, row_number, column_number)
                checked += 1
        results.append({"file": str(csv_path.relative_to(ROOT)), "cells": checked, "sha256": hashlib.sha256(csv_path.read_bytes()).hexdigest(), "result": "passed"})
print(json.dumps({"result": "passed", "files": len(results), "cells": sum(item["cells"] for item in results), "sheets": results}, ensure_ascii=False, indent=2))
