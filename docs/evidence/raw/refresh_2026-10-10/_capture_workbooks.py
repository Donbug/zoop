"""Capture public team XLSX exports; immutable baselines, no credentials or installs."""
import csv
import hashlib
import io
import json
import re
import sys
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
SOURCES = [
    ("wbs", "wbs_2026-10-08-api-review", "1aHJbxjXmvzgTXAlBd4mWTRDkcvD1S9vtHl1PR_bxn5c"),
    ("functional-spec", "functional-spec_2026-10-08", "1irLUY-NCkDRfF8fO3WkZWU2MpaStu2RPmQ-wG0F9Xcs"),
    ("requirements-sheet", "requirements-sheet_2026-10-06", "1zehuZhPlDNoC3g7EMNqlzmzUPyb9kiu31j0wO0udG4k"),
]

def sha(data):
    return hashlib.sha256(data).hexdigest()

def rel(path):
    return str(path.relative_to(ROOT.parent.parent.parent))

def write_new(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as target:
        target.write(data)

def save_json(path, value):
    write_new(path, (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))

def strings(worksheet):
    return [["" if value is None else str(value) for value in row]
            for row in worksheet.iter_rows(values_only=True)]

def hyperlinks(workbook):
    return {(ws.title, cell.coordinate): {
        "target": cell.hyperlink.target, "location": cell.hyperlink.location,
        "display": cell.hyperlink.display, "tooltip": cell.hyperlink.tooltip}
        for ws in workbook.worksheets for row in ws for cell in row if cell.hyperlink}

def fetch(source):
    key, _, source_id = source
    url = f"https://docs.google.com/spreadsheets/d/{source_id}/export?format=xlsx"
    started = datetime.now(timezone.utc)
    request = urllib.request.Request(url, headers={"User-Agent": "Zoop-Public-Evidence-Recheck/1.0"})
    with urllib.request.urlopen(request, timeout=90) as response:
        data = response.read()
        status = response.status
        content_type = response.headers.get("Content-Type")
    completed = datetime.now(timezone.utc)
    if status != 200 or not data.startswith(b"PK"):
        raise RuntimeError(f"{key}: expected public XLSX HTTP 200, got {status}, {content_type}")
    return source, data, {
        "source_url": f"https://docs.google.com/spreadsheets/d/{source_id}", "export_url": url,
        "request_started_at_utc": started.isoformat(), "collected_at_utc": completed.isoformat(),
        "collected_at_kst": completed.astimezone(ZoneInfo("Asia/Seoul")).isoformat(),
        "auth": "none", "http_status": status, "content_type": content_type,
        "response_bytes": len(data), "response_sha256": sha(data),
    }

def verify_csv(xlsx_path, sheet_files):
    workbook = load_workbook(xlsx_path, data_only=True)
    if list(sheet_files) != workbook.sheetnames:
        raise AssertionError((str(xlsx_path), "sheet names/order"))
    records = []
    for worksheet in workbook.worksheets:
        csv_path = xlsx_path.parent / sheet_files[worksheet.title]
        with csv_path.open(encoding="utf-8-sig", newline="") as source:
            actual = list(csv.reader(source))
        if len(actual) != worksheet.max_row:
            raise AssertionError((str(csv_path), "row count"))
        for row_number in range(1, worksheet.max_row + 1):
            if len(actual[row_number - 1]) != worksheet.max_column:
                raise AssertionError((str(csv_path), row_number, "column count"))
            for column_number in range(1, worksheet.max_column + 1):
                value = worksheet.cell(row_number, column_number).value
                expected = "" if value is None else str(value)
                if actual[row_number - 1][column_number - 1] != expected:
                    raise AssertionError((str(csv_path), row_number, column_number))
        records.append({"sheet": worksheet.title, "csv": rel(csv_path),
                        "rows": worksheet.max_row, "columns": worksheet.max_column,
                        "cells": worksheet.max_row * worksheet.max_column,
                        "sha256": sha(csv_path.read_bytes()), "result": "passed"})
    workbook.close()
    return records

baseline_hashes = {path: sha(path.read_bytes()) for _, name, _ in SOURCES
                   for path in (ROOT / name).iterdir() if path.is_file()}
with ThreadPoolExecutor(max_workers=3) as pool:
    fetched = list(pool.map(fetch, SOURCES))

all_deltas, all_formula_deltas, all_mapping_deltas = [], [], []
results, verification, relevant_rows = [], [], []
for (key, baseline_name, _), data, metadata in fetched:
    baseline_dir = ROOT / baseline_name
    baseline_manifest = json.loads((baseline_dir / "_collection-manifest.json").read_text())
    baseline_path = baseline_dir / baseline_manifest["xlsx"]
    old = load_workbook(baseline_path, data_only=True)
    new = load_workbook(io.BytesIO(data), data_only=True)
    old_raw = load_workbook(baseline_path, data_only=False)
    new_raw = load_workbook(io.BytesIO(data), data_only=False)
    before_links, after_links = hyperlinks(old_raw), hyperlinks(new_raw)
    hyperlink_deltas = [{"workbook": key, "sheet": sheet, "cell": cell,
                        "before": before_links.get((sheet, cell)), "after": after_links.get((sheet, cell))}
                       for sheet, cell in sorted(set(before_links) | set(after_links))
                       if before_links.get((sheet, cell)) != after_links.get((sheet, cell))]
    deltas, formula_deltas, structures, mapping_deltas = [], [], [], []
    for sheet in dict.fromkeys(old.sheetnames + new.sheetnames):
        bw = old[sheet] if sheet in old.sheetnames else None
        aw = new[sheet] if sheet in new.sheetnames else None
        before = strings(bw) if bw else []
        after = strings(aw) if aw else []
        before_shape = [bw.max_row, bw.max_column] if bw else None
        after_shape = [aw.max_row, aw.max_column] if aw else None
        if before_shape != after_shape:
            structures.append({"sheet": sheet, "before": before_shape, "after": after_shape})
        changed_rows = set()
        for row in range(1, max(len(before), len(after)) + 1):
            b = before[row - 1] if row <= len(before) else []
            a = after[row - 1] if row <= len(after) else []
            for column in range(1, max(len(b), len(a)) + 1):
                bv = b[column - 1] if column <= len(b) else ""
                av = a[column - 1] if column <= len(a) else ""
                if bv != av:
                    deltas.append({"workbook": key, "sheet": sheet, "row": row, "column": column,
                                   "cell": f"{get_column_letter(column)}{row}", "before": bv, "after": av})
                    changed_rows.add(row)
        br = old_raw[sheet] if sheet in old_raw.sheetnames else None
        ar = new_raw[sheet] if sheet in new_raw.sheetnames else None
        formulas_before = {cell.coordinate: str(cell.value) for row in br for cell in row if cell.data_type == "f"} if br else {}
        formulas_after = {cell.coordinate: str(cell.value) for row in ar for cell in row if cell.data_type == "f"} if ar else {}
        for cell in sorted(set(formulas_before) | set(formulas_after)):
            bv, av = formulas_before.get(cell, ""), formulas_after.get(cell, "")
            if bv != av:
                formula_deltas.append({"workbook": key, "sheet": sheet, "cell": cell, "before": bv, "after": av})
        for row in sorted(changed_rows):
            b = before[row - 1] if row <= len(before) else []
            a = after[row - 1] if row <= len(after) else []
            row_changes = [item for item in deltas if item["sheet"] == sheet and item["row"] == row]
            fs_matches = sorted(set(re.findall(r"\b(?:N?FS)\d+(?:-\d+)?\b", "\n".join(b + a))))
            if fs_matches or sheet == "명세 목차" or (key == "wbs" and sheet == "프로젝트 WBS"):
                mapping_deltas.append({"workbook": key, "sheet": sheet, "row": row,
                    "fs_nfs_ids_in_row": fs_matches,
                    "before_nonempty_cells": {get_column_letter(i + 1): v for i, v in enumerate(b) if v},
                    "after_nonempty_cells": {get_column_letter(i + 1): v for i, v in enumerate(a) if v},
                    "cell_deltas": row_changes})
        for row, values in enumerate(after, 1):
            text = "\n".join(values)
            if (key == "wbs" and "윤서진" in text) or (key == "functional-spec" and (sheet in {"명세 목차", "비기능 상세", "AI 챗봇", "AI 챗봇 v2"})) or (key == "requirements-sheet" and re.search("재무|추천|동시|지연|오류|부하", text)):
                relevant_rows.append({"workbook": key, "sheet": sheet, "row": row,
                                      "nonempty_cells": {get_column_letter(i + 1): v for i, v in enumerate(values) if v}})
    changed = bool(deltas or formula_deltas or hyperlink_deltas or structures or old.sheetnames != new.sheetnames)
    capture_dir = ROOT / f"{key}_2026-10-10" if changed else OUT / key
    xlsx_name = f"zoop-{key}_2026-10-10.xlsx"
    capture_path = capture_dir / xlsx_name
    write_new(capture_path, data)
    sheet_files = {}
    for index, worksheet in enumerate(new.worksheets):
        filename = baseline_manifest["sheets"].get(worksheet.title)
        if filename is None:
            filename = f"{index:02d}-" + re.sub(r"[^\w가-힣.-]+", "-", worksheet.title).strip("-") + ".csv"
        if filename in sheet_files.values():
            raise RuntimeError(f"CSV filename collision: {worksheet.title}: {filename}")
        sheet_files[worksheet.title] = filename
        stream = io.StringIO(newline="")
        csv.writer(stream).writerows(strings(worksheet))
        write_new(capture_dir / filename, stream.getvalue().encode("utf-8-sig"))
    manifest = dict(metadata, baseline=baseline_name, baseline_xlsx_sha256=sha(baseline_path.read_bytes()),
        content_identical=not changed, meaningful_cell_content_changed=changed,
        data_only=True, csv_encoding="utf-8-sig", xlsx=xlsx_name, sha256=sha(data), sheets=sheet_files,
        sheet_ranges={ws.title: {"rows": ws.max_row, "columns": ws.max_column} for ws in new.worksheets},
        files={p.name: sha(p.read_bytes()) for p in capture_dir.iterdir() if p.is_file()},
        baseline_sheet_order=old.sheetnames, current_sheet_order=new.sheetnames,
        cell_delta_count=len(deltas), formula_delta_count=len(formula_deltas), structural_delta_count=len(structures))
    save_json(capture_dir / "_collection-manifest.json", manifest)
    save_json(capture_dir / "_cell-delta.json", deltas)
    save_json(capture_dir / "_formula-delta.json", formula_deltas)
    save_json(capture_dir / "_structure-delta.json", structures)
    save_json(capture_dir / "_mapping-relevant-delta.json", mapping_deltas)
    save_json(capture_dir / "_hyperlink-targets.json", {
        "baseline": [{"sheet": s, "cell": c, **v} for (s, c), v in before_links.items()],
        "current": [{"sheet": s, "cell": c, **v} for (s, c), v in after_links.items()], "deltas": hyperlink_deltas})
    all_deltas.extend(deltas)
    all_formula_deltas.extend(formula_deltas)
    all_mapping_deltas.extend(mapping_deltas)
    current_verification = verify_csv(capture_path, sheet_files)
    baseline_verification = verify_csv(baseline_path, baseline_manifest["sheets"])
    verification.extend(current_verification)
    results.append(dict(metadata, workbook=key, baseline=rel(baseline_dir), capture=rel(capture_dir),
        recommended_importer_input=rel(capture_dir if changed else baseline_dir), content_identical=not changed,
        baseline_xlsx_sha256=sha(baseline_path.read_bytes()), sheet_count=len(new.sheetnames),
        cell_delta_count=len(deltas), formula_delta_count=len(formula_deltas), structural_deltas=structures,
        mapping_relevant_row_count=len(mapping_deltas), hyperlink_delta_count=len(hyperlink_deltas),
        current_hyperlink_count=len(after_links), cells_checked=sum(v["cells"] for v in current_verification),
        baseline_cells_checked=sum(v["cells"] for v in baseline_verification),
        baseline_csv_readback="passed", current_csv_readback="passed"))
    for workbook in (old, new, old_raw, new_raw):
        workbook.close()
unchanged_originals = all(sha(p.read_bytes()) == digest for p, digest in baseline_hashes.items())
if not unchanged_originals:
    raise AssertionError("original frozen file bytes changed")
save_json(OUT / "_cell-delta.json", all_deltas)
save_json(OUT / "_formula-delta.json", all_formula_deltas)
save_json(OUT / "_mapping-relevant-delta.json", all_mapping_deltas)
save_json(OUT / "_relevant-source-rows.json", relevant_rows)
summary = {"result": "passed", "method": {"python": sys.version, "openpyxl": openpyxl.__version__,
    "data_only": True, "csv_encoding": "utf-8-sig", "empty_cell": "", "other_values": "str(value)",
    "all_worksheets": True, "comparison": "exact string cell values, sheet names/order and rectangular ranges; formula and hyperlink target deltas recorded",
    "network": "three unauthenticated public XLSX GET requests; no external writes; no package installs"},
    "workbooks": results, "counts": {"workbooks": len(results), "sheets": len(verification),
    "current_cells_checked": sum(v["cells"] for v in verification), "cell_deltas": len(all_deltas),
    "formula_deltas": len(all_formula_deltas), "mapping_relevant_rows": len(all_mapping_deltas)},
    "invariants": {"all_tabs_captured": True, "full_workbook_to_csv_readback": "passed",
    "original_frozen_files_unchanged": unchanged_originals, "original_files_hash_checked": len(baseline_hashes)},
    "recommended_importer_inputs": {v["workbook"]: v["recommended_importer_input"] for v in results}}
save_json(OUT / "_source-recheck.json", summary)
save_json(OUT / "_csv-verification.json", {"result": "passed", "sheets": verification,
    "cells": sum(v["cells"] for v in verification)})
print(json.dumps(summary, ensure_ascii=False, indent=2))
